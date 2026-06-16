from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "chapter6_core_test_cases_35.md"
MD_OUTPUT = ROOT / "docs" / "chapter5_1_2_test_cases_by_module.md"
DOCX_OUTPUT = ROOT / "docs" / "chapter5_1_2_test_cases_by_module.docx"


@dataclass(frozen=True)
class Module:
    title: str
    intro: str
    case_numbers: tuple[int, ...]


MODULES = [
    Module(
        "5.1.2.1 登录注册模块的测试过程及结果",
        "登录注册模块主要验证默认账号登录、错误密码拦截、用户名和手机号唯一性、手机号格式校验以及个人资料手机号修改校验。",
        (1, 2, 3, 4, 6, 7),
    ),
    Module(
        "5.1.2.2 角色权限模块的测试过程及结果",
        "角色权限模块主要验证公开注册角色限制，以及普通用户只能操作本人预约和本人商品订单的数据归属边界。",
        (5, 17, 30),
    ),
    Module(
        "5.1.2.3 服务浏览与收藏模块的测试过程及结果",
        "服务浏览与收藏模块主要验证服务分类、服务关键字筛选和用户收藏状态是否与数据库服务数据保持一致。",
        (8, 9, 10),
    ),
    Module(
        "5.1.2.4 服务预约与技师推荐模块的测试过程及结果",
        "服务预约与技师推荐模块主要验证预约时间、技师服务能力、排班、预约冲突和自动推荐技师等核心规则。",
        (11, 12, 13, 14, 15),
    ),
    Module(
        "5.1.2.5 预约状态流转与评价模块的测试过程及结果",
        "预约状态流转与评价模块主要验证预约状态机、未完成预约评价拦截以及已评价预约重复评价拦截。",
        (16, 18, 19),
    ),
    Module(
        "5.1.2.6 请假管理模块的测试过程及结果",
        "请假管理模块主要验证技师请假提交、请假时间合法性、管理员驳回备注和请假审批预约冲突校验。",
        (20, 21, 22, 23),
    ),
    Module(
        "5.1.2.7 考勤打卡模块的测试过程及结果",
        "考勤打卡模块主要验证技师请假期间签到限制和已签到技师正常签退流程。",
        (24, 25),
    ),
    Module(
        "5.1.2.8 商品商城与购物车模块的测试过程及结果",
        "商品商城与购物车模块主要验证购物车库存限制、购物车金额计算和空商品订单提交拦截。",
        (26, 27, 28),
    ),
    Module(
        "5.1.2.9 商品订单与售后模块的测试过程及结果",
        "商品订单与售后模块主要验证库存不足下单、订单状态流转、售后申请和售后通过后的库存恢复。",
        (29, 31, 32, 33),
    ),
    Module(
        "5.1.2.10 钱包充值模块的测试过程及结果",
        "钱包充值模块主要验证充值活动支付成功后钱包余额、累计充值、累计赠送和交易流水是否正确。",
        (34,),
    ),
    Module(
        "5.1.2.11 AI顾问模块的测试过程及结果",
        "AI顾问模块主要验证无文本、无图片的无效咨询请求能否被后端正确拦截。",
        (35,),
    ),
]


def parse_cases() -> dict[int, tuple[str, list[tuple[str, str]]]]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    cases: dict[int, tuple[str, list[tuple[str, str]]]] = {}
    i = 0
    while i < len(lines):
        title = lines[i].strip()
        match = re.match(r"^表6-(\d+)\s+(.+)$", title)
        if not match:
            i += 1
            continue
        old_no = int(match.group(1))
        rows: list[tuple[str, str]] = []
        i += 1
        while i < len(lines):
            line = lines[i].strip()
            if re.match(r"^表6-\d+\s+", line):
                break
            if line.startswith("|") and not re.match(r"^\|\s*-+\s*\|", line):
                parts = [part.strip() for part in line.strip("|").split("|")]
                if len(parts) == 2 and parts[0] != "项目":
                    rows.append((parts[0], parts[1]))
            i += 1
        cases[old_no] = (match.group(2), rows)
    if len(cases) != 35:
        raise RuntimeError(f"Expected 35 test cases, got {len(cases)}")
    return cases


def table_markdown(table_no: int, title: str, rows: list[tuple[str, str]]) -> str:
    out = [f"表5-{table_no} {title}", "", "| 项目 | 内容 |", "| --- | --- |"]
    out.extend(f"| {key} | {value} |" for key, value in rows)
    out.append("")
    return "\n".join(out)


def build_markdown(cases: dict[int, tuple[str, list[tuple[str, str]]]]) -> None:
    table_no = 1
    parts = [
        "# 5.1.2 测试过程及结果",
        "",
        "本节按照系统功能模块整合核心测试用例。所有用例均依据项目后端代码、前端校验逻辑、数据库脚本 `bfsm.sql` 以及已通过的后端单元测试结果整理，测试数据采用系统真实账号、服务、商品、预约、订单、钱包和通知数据。",
        "",
    ]
    for module in MODULES:
        parts.extend([f"## {module.title}", "", module.intro, ""])
        for case_no in module.case_numbers:
            case_title, rows = cases[case_no]
            parts.append(table_markdown(table_no, case_title, rows))
            table_no += 1
    MD_OUTPUT.write_text("\n".join(parts).rstrip() + "\n", encoding="utf-8")


def set_run_font(run, east_asia: str = "宋体", latin: str = "Times New Roman", size: float | None = None) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 130, bottom: int = 90, end: int = 130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Cm(width_cm)


def configure_doc(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(0)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6 if level > 1 else 0)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, "黑体", size=16 if level == 1 else 12)
    run.bold = True


def add_paragraph(doc: Document, text: str, first_line: bool = True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(4)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(21)
    run = p.add_run(text)
    set_run_font(run, "宋体", size=10.5)


def add_case_table(doc: Document, caption_no: int, title: str, rows: list[tuple[str, str]]) -> None:
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(5)
    run = caption.add_run(f"表5-{caption_no} {title}")
    set_run_font(run, "黑体", size=10.5)
    run.bold = True

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for key, value in rows:
        cells = table.add_row().cells
        set_cell_width(cells[0], 3.1)
        set_cell_width(cells[1], 12.8)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        set_cell_shading(cells[0], "F2F5F8")

        lp = cells[0].paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.paragraph_format.line_spacing = 1.15
        lr = lp.add_run(key)
        set_run_font(lr, "黑体", size=9.5)
        lr.bold = True

        rp = cells[1].paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rp.paragraph_format.line_spacing = 1.2
        rr = rp.add_run(value)
        set_run_font(rr, "宋体", size=9.5)
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(0)


def build_docx(cases: dict[int, tuple[str, list[tuple[str, str]]]]) -> None:
    doc = Document()
    configure_doc(doc)
    add_heading(doc, "5.1.2 测试过程及结果", 1)
    add_paragraph(
        doc,
        "本节按照系统功能模块整合核心测试用例。所有用例均依据项目后端代码、前端校验逻辑、数据库脚本 bfsm.sql 以及已通过的后端单元测试结果整理，测试数据采用系统真实账号、服务、商品、预约、订单、钱包和通知数据。",
    )
    table_no = 1
    for module_index, module in enumerate(MODULES):
        add_heading(doc, module.title, 2)
        add_paragraph(doc, module.intro)
        for case_index, case_no in enumerate(module.case_numbers):
            case_title, rows = cases[case_no]
            add_case_table(doc, table_no, case_title, rows)
            table_no += 1
            is_last_case_in_module = case_index == len(module.case_numbers) - 1
            is_last_module = module_index == len(MODULES) - 1
            if not (is_last_case_in_module and is_last_module):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.core_properties.title = "5.1.2 测试过程及结果"
    doc.core_properties.subject = "美容美发预约管理系统核心测试用例模块整合"
    doc.core_properties.author = "Codex"
    doc.save(DOCX_OUTPUT)


def main() -> None:
    cases = parse_cases()
    build_markdown(cases)
    build_docx(cases)
    print(MD_OUTPUT)
    print(DOCX_OUTPUT)


if __name__ == "__main__":
    main()
