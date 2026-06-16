from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DOCX = ROOT / "docs" / "BFSM核心代码实现说明_增强注释版.docx"


SECTIONS = [
    {
        "title": "1. 预约创建与技师档期锁",
        "point": "该方法完成用户提交预约的核心流程，包含身份获取、服务时长计算、技师自动推荐、Redis 档期锁、二次冲突校验以及预约单落库。",
        "code": r'''@Transactional
public AppointmentDetailVo create(AppointmentCreateReq req) {
    // 注释：从 Spring Security 上下文中获取当前登录用户，保证预约单归属于真实用户。
    SecurityUser currentUser = SecurityUtils.currentUserOrThrow();

    // 注释：根据用户选择的服务项目读取服务价格、服务时长等快照信息。
    ServiceItem item = serviceItemOrThrow(req.getServiceItemId());
    // 注释：结束时间由开始时间 + 服务时长计算，后续冲突判断都使用完整时间区间。
    LocalDateTime appointmentEndTime =
            resolveAppointmentEndTime(item, req.getAppointmentTime());

    Long technicianId = req.getTechnicianId();
    String assignMode = "MANUAL";
    if (technicianId == null) {
        // 注释：用户未手动选择技师时，系统根据排班、请假、冲突和评分自动推荐技师。
        TechnicianRecommendationVo recommended =
                schedulingService.recommendTechnician(req.getServiceItemId(), req.getAppointmentTime());
        if (recommended == null) {
            throw new BizException(4011, "当前时间段没有可推荐的技师");
        }
        technicianId = recommended.getTechnicianId();
        assignMode = "AUTO";
    }

    // 注释：以“技师 + 日期”为维度加锁，防止多个用户同时抢占同一技师档期。
    String lockKey = lockKey(technicianId, req.getAppointmentTime());
    String lockToken = redisLockService.tryLock(lockKey, Duration.ofSeconds(8));
    if (lockToken == null) {
        throw new BizException(4012, "预约请求过于频繁，请稍后再试");
    }

    try {
        // 注释：获得锁后再次校验，避免加锁前后档期状态发生变化。
        scheduleConflictChecker.validateTechnicianForAppointment(
                technicianId, req.getServiceItemId(), req.getAppointmentTime(), appointmentEndTime);
        // 注释：检测该技师在预约区间内是否已有待确认、已确认、待服务或服务中的预约。
        if (scheduleConflictChecker.hasConflictInRange(
                technicianId, req.getAppointmentTime(), appointmentEndTime)) {
            throw new BizException(4013, "该时间段技师已被预约");
        }

        // 注释：保存预约主表，初始状态为 CREATED，支付状态为 UNPAID。
        Appointment appointment = new Appointment();
        appointment.setUserId(currentUser.getUserId());
        appointment.setTechnicianId(technicianId);
        appointment.setServiceItemId(req.getServiceItemId());
        appointment.setAppointmentTime(req.getAppointmentTime());
        appointment.setAppointmentEndTime(appointmentEndTime);
        appointment.setStatus("CREATED");
        appointment.setPayStatus("UNPAID");
        appointment.setAmount(item.getPrice());
        appointment.setAssignMode(assignMode);
        appointmentMapper.insert(appointment);
        return detail(appointment.getId());
    } finally {
        // 注释：finally 中释放锁，避免业务异常导致锁长时间占用。
        redisLockService.unlock(lockKey, lockToken);
    }
}''',
    },
    {
        "title": "2. 技师可用性与冲突检查",
        "point": "该模块将技师是否可预约拆成三个条件：能否提供该服务、是否在排班时间内、是否处于请假时间段，并单独提供预约时间重叠检测。",
        "code": r'''public void validateTechnicianForAppointment(Long technicianId, Long serviceItemId,
                                             LocalDateTime startTime, LocalDateTime endTime) {
    // 注释：校验技师是否绑定了当前服务项目，避免用户预约到不具备服务能力的技师。
    if (!supportsService(technicianId, serviceItemId)) {
        throw new BizException(4008, "Selected technician does not support this service item");
    }
    // 注释：校验预约时间是否完整落在技师当天排班时间内。
    if (!isWorkingDuring(technicianId, startTime, endTime)) {
        throw new BizException(4009, "Selected technician is not working in this time range");
    }
    // 注释：校验技师在该时间段是否存在已批准请假记录。
    if (isOnLeaveDuring(technicianId, startTime, endTime)) {
        throw new BizException(4311, "Selected technician is on leave in this period");
    }
}

public boolean hasConflictInRange(Long technicianId, LocalDateTime startTime, LocalDateTime endTime) {
    // 注释：通过数据库查询判断时间区间是否重叠，是预约防冲突的核心判断。
    return appointmentMapper.countConflict(technicianId, startTime, endTime) > 0;
}''',
    },
    {
        "title": "3. 自动推荐技师",
        "point": "推荐逻辑先过滤不可服务技师，再按照评分和当天负载计算综合得分，实现用户无指定技师时的自动分配。",
        "code": r'''public TechnicianRecommendationVo recommendFromCandidates(ServiceItem item,
        LocalDateTime appointmentTime, List<Technician> candidates) {
    // 注释：根据服务时长计算预约结束时间，确保推荐结果能覆盖完整服务过程。
    LocalDateTime appointmentEndTime = appointmentTime.plusMinutes(item.getDurationMin());

    return candidates.stream()
            // 注释：只保留该时间段正在排班的技师。
            .filter(t -> scheduleConflictChecker.isWorkingDuring(t.getId(), appointmentTime, appointmentEndTime))
            // 注释：排除请假中的技师。
            .filter(t -> !scheduleConflictChecker.isOnLeaveDuring(t.getId(), appointmentTime, appointmentEndTime))
            // 注释：排除已有预约冲突的技师。
            .filter(t -> !scheduleConflictChecker.hasConflictInRange(t.getId(), appointmentTime, appointmentEndTime))
            // 注释：将可服务技师转换为推荐结果，并计算推荐分数。
            .map(t -> toRecommendation(t, appointmentTime, appointmentEndTime, item.getDurationMin()))
            // 注释：选择综合得分最高的技师作为最终推荐结果。
            .max(Comparator.comparingDouble(TechnicianRecommendationVo::getScore))
            .orElse(null);
}

private TechnicianRecommendationVo toRecommendation(Technician technician,
        LocalDateTime appointmentTime, LocalDateTime appointmentEndTime, Integer durationMin) {
    // 注释：统计技师当天已有预约数量，作为负载因子。
    long dayLoad = appointmentMapper.countDayLoad(technician.getId(), appointmentTime.toLocalDate());
    // 注释：没有历史评分时使用默认评分 5.0，避免空值影响推荐计算。
    double rating = technician.getScore() == null ? 5.0D : technician.getScore().doubleValue();
    // 注释：推荐分数 = 基础分 + 评分加权 - 当天负载惩罚，兼顾服务质量和工作均衡。
    double score = 60 + rating * 8 - dayLoad * 10;
    ...
}''',
    },
    {
        "title": "4. 商品下单与库存并发控制",
        "point": "该方法负责商品订单创建，重点是通过 Redis 库存锁防止并发扣减导致超卖，同时保存订单、订单明细和收货地址快照。",
        "code": r'''@Transactional
public ProductOrderDetailVo createOrder(ProductOrderCreateReq req) {
    // 注释：获取当前用户，确保用户只能为自己的购物车或商品创建订单。
    SecurityUser currentUser = SecurityUtils.currentUserOrThrow();
    // 注释：读取收货地址并保存为订单快照，避免用户后续修改地址影响历史订单。
    UserAddress address = currentUserAddress(currentUser.getUserId(), req.getAddressId());
    // 注释：解析订单明细，兼容购物车结算和直接购买两种来源。
    List<OrderLine> lines = resolveOrderLines(currentUser.getUserId(), req);

    Map<String, String> lockTokens = new LinkedHashMap<>();
    Map<Long, Product> products = new LinkedHashMap<>();
    try {
        for (OrderLine line : lines) {
            Product product = productMapper.selectById(line.productId());
            // 注释：按商品 ID 加库存锁，保证同一商品的库存扣减串行执行。
            String lockKey = "lock:product:stock:" + product.getId();
            String token = redisLockService.tryLock(lockKey, Duration.ofSeconds(8));
            lockTokens.put(lockKey, token);

            // 注释：加锁后重新读取最新库存，避免使用旧库存数据。
            Product latest = productMapper.selectById(product.getId());
            if (latest.getStock() < line.quantity()) {
                throw new BizException(4102, "商品库存不足");
            }
            // 注释：库存充足时立即扣减，防止后续请求重复占用同一批库存。
            latest.setStock(latest.getStock() - line.quantity());
            productMapper.updateById(latest);
            products.put(latest.getId(), latest);
        }

        // 注释：创建订单主表，订单初始状态为 CREATED，支付状态为 UNPAID。
        ProductOrder order = new ProductOrder();
        order.setUserId(currentUser.getUserId());
        order.setStatus("CREATED");
        order.setPayStatus("UNPAID");
        order.setReceiverName(address.getReceiverName());
        order.setTotalAmount(total);
        productOrderMapper.insert(order);
        ...
        return detailForCurrentUser(order.getId());
    } finally {
        // 注释：无论下单成功或失败，都释放所有已获取的商品库存锁。
        lockTokens.forEach((key, token) -> redisLockService.unlock(key, token));
    }
}''',
    },
    {
        "title": "5. 统一支付单与业务回写",
        "point": "支付模块把预约、商品订单、钱包充值统一抽象成 PaymentOrder，支付渠道由网关路由处理，支付成功后再回写具体业务状态。",
        "code": r'''@Transactional
public PaymentCreateVo createAppointmentPayment(Long appointmentId, PaymentCreateReq req) {
    // 注释：获取当前用户，防止用户为他人的预约发起支付。
    SecurityUser currentUser = SecurityUtils.currentUserOrThrow();
    // 注释：规范化支付渠道，例如微信、支付宝、手动收银等。
    String payChannel = normalizeCheckoutChannel(req.getPayChannel());
    // 注释：根据渠道解析实际支付模式，用于路由到具体支付网关。
    String payMode = resolveSupportedPayMode(payChannel);

    // 注释：查询当前用户自己的预约单，并进行支付前状态校验。
    Appointment appointment = requireCurrentUserAppointment(appointmentId, currentUser.getUserId());
    if (!"CREATED".equals(appointment.getStatus())
            || !"UNPAID".equals(normalizeAppointmentPayStatus(appointment))) {
        throw new BizException(4711, "当前预约不可发起支付");
    }

    // 注释：同一预约重新发起支付时，关闭旧的待支付支付单，避免多支付单并存。
    closePendingPayments("APPOINTMENT", appointment.getId(), "A new payment was created");

    // 注释：创建统一支付单，业务类型标记为 APPOINTMENT。
    PaymentOrder paymentOrder = new PaymentOrder();
    paymentOrder.setBusinessType("APPOINTMENT");
    paymentOrder.setBusinessId(appointment.getId());
    paymentOrder.setPaymentNo(buildPaymentNo("A"));
    paymentOrder.setPayChannel(payChannel);
    paymentOrder.setPayMode(payMode);
    paymentOrder.setAmount(scale(appointment.getAmount()));
    paymentOrder.setStatus("PAYING");
    // 注释：设置 15 分钟支付有效期，过期后由定时任务关闭。
    paymentOrder.setExpireTime(LocalDateTime.now().plusMinutes(15));
    paymentOrderMapper.insert(paymentOrder);

    // 注释：根据支付模式进入对应网关，返回二维码、跳转表单或手动确认信息。
    return gatewayForMode(payMode).create(paymentOrder);
}''',
    },
    {
        "title": "6. AI 咨询推荐闭环",
        "point": "AI 咨询不是独立问答，而是把用户问题、图片、服务目录和商品目录组合成上下文，优先调用远程 AI，失败时使用本地规则兜底。",
        "code": r'''private TurnResult buildTurn(String question, String imageUrl,
                             AiConsultRecord record, Catalog catalog) {
    // 注释：识别用户意图，判断本轮是发型建议、服务推荐、商品推荐还是普通问答。
    Intent intent = detectIntent(question);
    // 注释：优先调用远程 AI 接口；接口不可用时可返回 null，后续走本地推荐规则。
    AiBeautyAnalysis remote = tryRemoteAnalysis(question, imageUrl, record, catalog);

    // 注释：将远程 AI 返回的服务推荐与系统真实服务目录合并，过滤不存在或已下架的服务。
    List<AiConsultServiceRecommendationVo> services =
            mergeServiceRecommendations(
                    remote == null ? List.of() : remote.getServiceRecommendations(),
                    catalog.servicesById());

    // 注释：将远程 AI 返回的商品推荐与真实商品目录合并，保证推荐结果可购买。
    List<AiConsultProductRecommendationVo> products =
            mergeProductRecommendations(
                    remote == null ? List.of() : remote.getProductRecommendations(),
                    catalog.productsById());

    // 注释：当远程 AI 未返回服务推荐时，根据关键词和服务目录做本地兜底推荐。
    if (intent.serviceIntent() && services.isEmpty()) {
        services = recommendServices(question, catalog.services());
    }
    // 注释：当远程 AI 未返回商品推荐时，根据关键词和商品目录做本地兜底推荐。
    if (intent.productIntent() && products.isEmpty()) {
        products = recommendProducts(question, catalog.products());
    }

    // 注释：记录本轮分析模式，便于前端展示和后续问题排查。
    String mode = remote == null ? "LOCAL_RULE" : "REMOTE_API";
    return new TurnResult(stage, summary, observation, assistant,
            hairStyles, services, products, careTips, tags, mode, raw);
}''',
    },
]


def set_run_font(run, name: str, east_asia: str = "Microsoft YaHei", size: Pt | None = None):
    run.font.name = name
    if size is not None:
        run.font.size = size
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)


def shade(paragraph, fill: str):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def configure(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    code = doc.styles.add_style("AnnotatedCode", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.04)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("BFSM 核心代码实现说明（增强注释版）")
    set_run_font(r, "Calibri", size=Pt(18))
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("1F4D78")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("围绕预约、排班、库存、支付和 AI 推荐等核心业务流程补充关键注释")
    set_run_font(r, "Calibri", size=Pt(11))
    r.font.color.rgb = RGBColor(100, 100, 100)


def add_section(doc: Document, item: dict):
    doc.add_heading(item["title"], level=1)

    p = doc.add_paragraph()
    shade(p, "F4F6F9")
    r = p.add_run("实现要点：" + item["point"])
    set_run_font(r, "Calibri", size=Pt(10))

    label = doc.add_paragraph()
    r = label.add_run("核心代码：")
    set_run_font(r, "Calibri", size=Pt(10))
    r.bold = True

    for line in item["code"].splitlines():
        p = doc.add_paragraph(style="AnnotatedCode")
        shade(p, "F7F8FA")
        r = p.add_run(line if line else " ")
        set_run_font(r, "Consolas", "SimSun", Pt(8))


def build():
    doc = Document()
    configure(doc)
    add_title(doc)
    for i, section in enumerate(SECTIONS):
        if i == 3:
            doc.add_section(WD_SECTION.NEW_PAGE)
        add_section(doc, section)
    doc.core_properties.title = "BFSM 核心代码实现说明（增强注释版）"
    doc.core_properties.subject = "核心代码摘录与必要注释"
    doc.core_properties.author = "Codex"
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build()
