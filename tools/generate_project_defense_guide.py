from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "基于Spring Boot的美容美发预约管理系统_项目答辩学习手册.docx"

TITLE = "基于 Spring Boot 的美容美发预约管理系统"
SUBTITLE = "项目答辩学习手册：从项目结构、核心代码到现场问答"
AUTHOR = "李洪东（学号：220162704018）"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202733"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "227447"
GOLD = "7A5A00"
RED = "9B1C1C"

TABLE_WIDTH = 9360
TABLE_INDENT = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"Table widths must total {TABLE_WIDTH}: {widths}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, latin: str = "Calibri",
                 east_asia: str = "Microsoft YaHei", italic: bool | None = None) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, color: str = "D0D5DD", size: str = "4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:space"), "4")
        edge.set(qn("w:color"), color)
        p_bdr.append(edge)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code_style = styles.add_style("CodeBlock", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code_style.font.size = Pt(7.5)
    code_style.font.color.rgb = RGBColor.from_string("1F2937")
    code_style.paragraph_format.left_indent = Inches(0.12)
    code_style.paragraph_format.right_indent = Inches(0.08)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(7)
    code_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for sec in doc.sections:
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        hr = hp.add_run("BFSM 项目答辩学习手册")
        set_run_font(hr, 8.5, True, MUTED)

        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_after = Pt(0)
        fr = fp.add_run("第 ")
        set_run_font(fr, 8.5, color=MUTED)
        add_page_number(fp)
        fr2 = fp.add_run(" 页")
        set_run_font(fr2, 8.5, color=MUTED)


def add_para(doc: Document, text: str = "", *, bold_prefix: str | None = None,
             color: str | None = None, align=None, after: float = 6,
             italic: bool = False) -> object:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, 10.5, True, color or INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, 10.5, False, color or INK, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, 10.5, False, color or INK, italic=italic)
    return p


def add_bullet(doc: Document, text: str) -> object:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, 10.5, color=INK)
    return p


def add_number(doc: Document, text: str) -> object:
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    set_run_font(run, 10.5, color=INK)
    return p


def add_heading(doc: Document, text: str, level: int = 1) -> object:
    return doc.add_heading(text, level=level)


def add_callout(doc: Document, label: str, text: str, kind: str = "info") -> object:
    colors = {
        "info": (CALLOUT, BLUE),
        "good": ("EAF5EE", GREEN),
        "warn": ("FFF7E6", GOLD),
        "risk": ("FDECEC", RED),
    }
    fill, accent = colors[kind]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.06)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_shading(p, fill)
    set_paragraph_border(p, color=accent, size="5")
    r1 = p.add_run(label + " ")
    set_run_font(r1, 10.5, True, accent)
    r2 = p.add_run(text)
    set_run_font(r2, 10.5, False, INK)
    return p


def add_code(doc: Document, source: str, code: str, explanation: str | None = None) -> object:
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(2)
    r1 = cap.add_run("源码位置  ")
    set_run_font(r1, 8.5, True, BLUE)
    r2 = cap.add_run(source)
    set_run_font(r2, 8.5, False, MUTED, latin="Consolas")

    p = doc.add_paragraph(style="CodeBlock")
    set_paragraph_shading(p, "F7F8FA")
    set_paragraph_border(p, color="D0D5DD", size="3")
    lines = code.strip("\n").splitlines()
    for idx, line in enumerate(lines):
        run = p.add_run(line.rstrip())
        set_run_font(run, 7.5, False, "1F2937", latin="Consolas", east_asia="Microsoft YaHei")
        if idx < len(lines) - 1:
            run.add_break()
    if explanation:
        add_callout(doc, "读代码时抓住：", explanation, "info")
    return p


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]],
              widths: list[int], header_fill: str = LIGHT_BLUE,
              center_cols: set[int] | None = None) -> object:
    rows = [list(map(str, row)) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                set_run_font(r, 9, True, NAVY)

    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center_cols and idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_run_font(r, 8.7, False, INK)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    kr = kicker.add_run("GRADUATION PROJECT FIELD GUIDE")
    set_run_font(kr, 10, True, BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(TITLE)
    set_run_font(r, 27, True, NAVY)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(42)
    r2 = p2.add_run(SUBTITLE)
    set_run_font(r2, 14, False, DARK_BLUE)

    add_callout(
        doc,
        "学习目标：",
        "在较短时间内建立对项目全貌、核心业务、关键代码、数据库关系和答辩表达的系统认识，做到能演示、能解释、能追问、能定位问题。",
        "good",
    )

    for _ in range(4):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(5)
    mr = meta.add_run(AUTHOR)
    set_run_font(mr, 11, True, INK)
    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta2.paragraph_format.space_after = Pt(0)
    mr2 = meta2.add_run("依据论文初稿与 E:\\Java\\bfsm 当前源码整理 | 2026 年 6 月")
    set_run_font(mr2, 9.5, False, MUTED)
    page_break(doc)


def add_contents(doc: Document) -> None:
    add_heading(doc, "使用说明与学习路线", 1)
    add_para(doc, "这份手册不是论文正文的重复，也不是源码清单。它按照答辩准备的认知顺序组织：先用几分钟讲清项目，再理解分层结构和请求链路，随后深入预约、支付、商城、排班、AI 等关键模块，最后通过问答与演示脚本完成复习。")
    add_table(
        doc,
        ["学习阶段", "建议时间", "需要达到的结果"],
        [
            ("第一遍：建立全貌", "30 分钟", "能用 2 至 3 分钟完整介绍项目，知道三类角色和三条核心业务线。"),
            ("第二遍：理解核心", "1 至 2 小时", "能讲清预约防冲突、状态机、统一支付、库存锁和权限控制。"),
            ("第三遍：对照源码", "2 至 3 小时", "能根据类名定位 Controller、Service、Mapper、前端页面和数据库表。"),
            ("第四遍：模拟答辩", "1 小时", "能回答高频问题，并按既定顺序完成 8 至 12 分钟演示。"),
        ],
        [1700, 1500, 6160],
        center_cols={1},
    )
    add_callout(doc, "最重要的原则：", "回答时先讲业务问题，再讲技术方案，最后指出代码位置和设计取舍。不要一上来只背框架名称。", "warn")

    add_heading(doc, "目录", 1)
    contents = [
        "1. 三分钟讲清项目",
        "2. 项目结构与系统架构",
        "3. 一次请求如何穿过前后端",
        "4. 登录认证与多角色权限",
        "5. 服务预约、时间段与技师推荐",
        "6. 排班、请假与考勤协同",
        "7. 统一支付、钱包与退款",
        "8. 商品商城、库存与售后",
        "9. AI 顾问、聊天、通知与统计",
        "10. 数据库设计与实体关系",
        "11. 代码阅读、调试与风险认识",
        "12. 高频答辩问题与参考回答",
        "13. 现场演示脚本与故障预案",
        "14. 最后一页速记",
    ]
    for item in contents:
        add_bullet(doc, item)
    page_break(doc)


def section_project_overview(doc: Document) -> None:
    add_heading(doc, "1. 三分钟讲清项目", 1)
    add_heading(doc, "1.1 项目定位", 2)
    add_para(doc, "本项目面向美容美发门店的线上化运营场景，采用 Spring Boot 与 Vue 3 构建前后端分离系统。系统不仅提供用户预约功能，还将技师排班、请假考勤、服务履约、商品销售、支付退款、消息沟通、AI 咨询和运营统计组织为相互关联的业务闭环。")
    add_callout(
        doc,
        "答辩开场可直接说：",
        "我的毕业设计是一个面向美容美发门店的预约与运营管理系统。系统采用 Vue 3 和 Spring Boot 前后端分离架构，以服务预约为核心，同时实现技师排班、支付结算、商品商城、AI 咨询、消息通知和运营统计。设计重点不是页面数量，而是处理预约时间冲突、库存并发扣减、多角色权限、订单状态流转以及支付成功后的业务回写。",
        "good",
    )

    add_heading(doc, "1.2 三类角色与业务范围", 2)
    add_table(
        doc,
        ["角色", "主要功能", "答辩中应强调的价值"],
        [
            ("普通用户 USER", "注册登录、浏览服务、预约技师、在线支付、钱包充值、商城购物、收藏评价、AI 咨询、聊天与通知", "完成从了解服务到预约、消费、评价的用户闭环"),
            ("技师 TECH", "处理预约状态、查看工作台、考勤打卡、提交请假、查看个人统计、与用户沟通", "把线上预约与线下履约连接起来"),
            ("管理员 ADMIN", "用户、服务、技师、排班、预约、请假、商品、订单、内容、评论、充值活动和统计管理", "维护平台基础数据并处理异常业务"),
        ],
        [1700, 4800, 2860],
        center_cols={0},
    )

    add_heading(doc, "1.3 项目规模快照", 2)
    add_para(doc, "按照当前工作区源码统计，后端包含 231 个主代码 Java 文件、17 个 Controller、39 个 Service、32 个 Entity、32 个 Mapper、31 个 DTO 和 52 个 VO；前端包含 44 个 Vue 单文件组件和 12 个 TypeScript 文件；主数据库脚本定义 31 张业务表，另有 AI 相关初始化器负责补充咨询记录表。")
    add_callout(doc, "如何解释项目规模：", "不要只说“功能很多”。应说明代码按领域拆分，预约、支付、商城、排班、消息等模块分别维护业务规则，并通过通知、支付和用户身份形成协作。", "info")

    add_heading(doc, "1.4 四个最值得讲的技术亮点", 2)
    for text in [
        "预约并发控制：以“技师 + 日期”为锁粒度，获取 Redis 锁后再次校验排班、请假和预约冲突，降低重复占用档期的风险。",
        "统一支付抽象：商品订单、服务预约和钱包充值共同使用 PaymentOrder，由支付网关路由选择支付宝、微信或手动收银模式，支付成功后统一回写业务。",
        "完整状态机：预约与商品订单都限制合法状态迁移，避免接口被任意调用后跳过业务步骤。",
        "AI 与真实目录融合：远程模型输出必须映射到系统中真实存在且可用的服务和商品，远程不可用时可按配置采用本地规则降级。",
    ]:
        add_bullet(doc, text)


def section_architecture(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "2. 项目结构与系统架构", 1)
    add_heading(doc, "2.1 顶层目录", 2)
    add_code(
        doc,
        "E:\\Java\\bfsm",
        r"""
bfsm/
├─ backend/                   Spring Boot 后端工程
│  ├─ src/main/java/com/bfsm
│  │  ├─ controller/          REST 接口入口
│  │  ├─ service/             业务规则与事务
│  │  ├─ mapper/              MyBatis-Plus 数据访问
│  │  ├─ entity/              数据库实体
│  │  ├─ dto/                 请求参数对象
│  │  ├─ vo/                  返回视图对象
│  │  ├─ security/            JWT 与 Spring Security
│  │  ├─ cache/               Redis 缓存与分布式锁
│  │  ├─ websocket/           实时聊天推送
│  │  └─ config/              跨域、支付、AI、静态资源等配置
│  └─ src/main/resources      application.yml、schema.sql
├─ frontend/                  Vue 3 + TypeScript 前端工程
│  └─ src/
│     ├─ views/               用户端、技师端、管理端页面
│     ├─ api/                 API 调用封装
│     ├─ router/              路由与角色守卫
│     ├─ store/               Pinia 全局状态
│     ├─ components/          公共组件
│     └─ utils/               Axios、校验、头像和分页工具
├─ docs/                      论文、图表和项目说明材料
└─ tools/                     文档及图表生成脚本
""",
        "先记住目录职责，不要试图记住全部文件。答辩被问到某个功能时，按照“页面 -> API -> Controller -> Service -> Mapper/表”的路径定位。",
    )

    add_heading(doc, "2.2 技术栈及其职责", 2)
    add_table(
        doc,
        ["技术", "在项目中的实际职责", "对应位置"],
        [
            ("Spring Boot 3.4.4 / Java 17", "提供 Web、依赖注入、事务、配置和定时任务等基础能力", "backend/pom.xml"),
            ("Spring Security + JWT + BCrypt", "无状态认证、方法级角色授权、密码哈希存储", "security/、AuthService"),
            ("MyBatis-Plus + MySQL", "实体映射、条件查询、分页和自定义 SQL", "entity/、mapper/、schema.sql"),
            ("Redis", "缓存当前用户信息；实现预约档期锁和商品库存锁", "cache/"),
            ("WebSocket", "聊天消息落库后向在线接收者实时推送", "websocket/、ChatService"),
            ("Vue 3 + TypeScript + Vite", "多角色前端页面、组件化开发与构建", "frontend/src"),
            ("Pinia + Vue Router + Axios", "登录状态、路由守卫、统一请求与错误处理", "store/、router/、utils/request.ts"),
            ("Element Plus + ECharts", "表格表单、弹窗分页和统计图表", "frontend/src/views"),
            ("支付网关抽象", "支付宝页面支付、微信 Native 支付、手动收银模式", "service/gateway/"),
            ("OpenAI-compatible 接口", "美容建议、服务和商品推荐，支持目录约束与降级", "service/ai/、AiConsultService"),
        ],
        [1900, 4800, 2660],
    )

    add_heading(doc, "2.3 系统分层", 2)
    add_code(
        doc,
        "全局调用关系",
        """
Vue 页面
  -> api/modules.ts 或 Axios request
  -> Spring Security / JwtAuthFilter
  -> Controller 接收 DTO
  -> Service 执行业务规则与事务
  -> Mapper 操作 MySQL
  -> Redis / 支付网关 / WebSocket / AI（按需）
  -> VO + ApiResponse 返回前端
""",
        "Controller 尽量保持轻量，复杂判断集中在 Service；Mapper 负责查询和持久化；DTO 与 VO 避免直接把数据库实体暴露给前端。",
    )

    add_heading(doc, "2.4 为什么采用前后端分离", 2)
    add_para(doc, "前后端分离使页面展示与业务规则解耦。Vue 负责不同角色的交互和数据展示，Spring Boot 通过 REST API 统一提供能力。JWT 适合这种无状态调用方式，前端每次请求携带令牌，后端独立验证身份。")
    add_callout(doc, "答辩追问：", "前后端分离不等于把权限交给前端。前端路由守卫只负责隐藏不应出现的页面，真正的接口安全由后端认证过滤器和 @PreAuthorize 控制。", "warn")


def section_request_flow(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "3. 一次请求如何穿过前后端", 1)
    add_heading(doc, "3.1 以“提交预约”为例", 2)
    for text in [
        "用户在 ServiceListView 或预约相关页面选择服务、日期、时间和技师。",
        "前端通过 Axios 向 POST /api/appointments 发送 AppointmentCreateReq。",
        "请求拦截器从 Pinia 登录状态读取 token，并添加 Authorization: Bearer <token>。",
        "JwtAuthFilter 解析令牌，把 SecurityUser 放入 Spring Security 上下文。",
        "AppointmentController 使用 @Valid 校验请求，然后调用 AppointmentService.create。",
        "Service 检查时间、服务项目、技师可用性和 Redis 锁，并在事务中写入 appointment 与 appointment_log。",
        "Controller 将 AppointmentDetailVo 包装为 ApiResponse 返回。",
        "前端响应拦截器检查 code，成功时刷新页面；失败时统一提示 message。",
    ]:
        add_number(doc, text)

    add_heading(doc, "3.2 统一响应与异常", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/common/ApiResponse.java",
        """
public class ApiResponse<T> {
    private int code;
    private String message;
    private T data;

    public static <T> ApiResponse<T> ok(T data) {
        return new ApiResponse<>(0, "OK", data);
    }
}
""",
        "前端不能只看 HTTP 状态，还会检查响应体 code。业务异常由 BizException 表达，GlobalExceptionHandler 将异常转换为统一结构。",
    )
    add_para(doc, "这种设计让所有页面使用相同的成功与失败处理逻辑，避免每个页面分别判断后端返回格式。需要注意的是，当前 GlobalExceptionHandler 将业务异常统一映射为 HTTP 400，因此前端主要依靠业务 code 和 message 区分具体原因。")

    add_heading(doc, "3.3 DTO、Entity 和 VO 的区别", 2)
    add_table(
        doc,
        ["类型", "含义", "项目示例", "为什么需要"],
        [
            ("DTO", "Data Transfer Object，请求输入", "AppointmentCreateReq、PaymentCreateReq", "限定前端可提交字段并配合参数校验"),
            ("Entity", "数据库实体", "Appointment、PaymentOrder、SysUser", "与数据库表字段对应，供 Mapper 持久化"),
            ("VO", "View Object，返回视图", "AppointmentDetailVo、PaymentCreateVo", "组合页面需要的信息，隐藏内部字段"),
        ],
        [1200, 2100, 2600, 3460],
        center_cols={0},
    )


def section_security(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "4. 登录认证与多角色权限", 1)
    add_heading(doc, "4.1 登录流程", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/AuthService.java",
        """
if (!passwordEncoder.matches(req.getPassword(), user.getPassword())) {
    throw new BizException(4005, "用户名或密码不正确");
}
SysRole role = roleMapper.selectById(userRole.getRoleId());
String token = jwtTokenService.createAccessToken(
        user.getId(), user.getUsername(), role.getCode());
return new AuthTokenVo(token, jwtProperties.getAccessTokenTtlSeconds());
""",
        "密码不是解密比较，而是使用 BCrypt 的 matches；JWT 中保存 userId、username 和 role，默认有效期由 application.yml 配置为 7200 秒。",
    )
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/security/JwtTokenService.java",
        """
return Jwts.builder()
        .setIssuer(props.getIssuer())
        .setSubject(String.valueOf(userId))
        .setIssuedAt(Date.from(now))
        .setExpiration(Date.from(exp))
        .addClaims(Map.of("username", username, "role", roleCode))
        .signWith(key, SignatureAlgorithm.HS256)
        .compact();
""",
        "令牌经过 HMAC-SHA256 签名，后端能检测内容是否被篡改。JWT 只是签名编码，不应在其中存放密码等敏感信息。",
    )

    add_heading(doc, "4.2 后端安全链路", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/security/SecurityConfig.java",
        """
http
    .csrf(csrf -> csrf.disable())
    .sessionManagement(sm ->
        sm.sessionCreationPolicy(SessionCreationPolicy.STATELESS))
    .authorizeHttpRequests(auth -> auth
        .requestMatchers("/api/auth/**", "/uploads/**", "/ws/**",
                         "/v3/api-docs/**", "/doc.html").permitAll()
        .anyRequest().authenticated());
http.addFilterBefore(jwtAuthFilter,
        UsernamePasswordAuthenticationFilter.class);
""",
        "系统不创建 Session；除登录、支付回调、静态资源和接口文档等白名单外，其余接口都要求认证。",
    )
    add_para(doc, "JwtAuthFilter 是 OncePerRequestFilter。它从 Authorization 请求头中提取 Bearer Token，解析出用户与角色后构造 UsernamePasswordAuthenticationToken，并写入 SecurityContextHolder。后续 Service 可以通过 SecurityUtils 获取当前用户，@PreAuthorize 也能判断角色。")

    add_heading(doc, "4.3 前端路由与后端方法授权", 2)
    add_table(
        doc,
        ["控制层次", "项目实现", "作用"],
        [
            ("前端", "router.beforeEach + route.meta.roles", "阻止用户进入不属于自身角色的页面，改善体验"),
            ("后端接口", "SecurityConfig 要求认证", "防止未登录请求访问业务接口"),
            ("后端方法", "@PreAuthorize(\"hasRole('ADMIN')\")", "限制管理员或技师才能执行敏感操作"),
            ("业务对象", "Service 中 assertCanView / assertCanOperate", "确保用户只能操作自己的订单或被分配的预约"),
        ],
        [1500, 3450, 4410],
    )
    add_callout(doc, "安全边界：", "即使用户手工构造 HTTP 请求绕过前端，后端仍会验证身份、角色和数据归属。前端权限不是安全措施的终点。", "risk")

    add_heading(doc, "4.4 可以主动承认的改进点", 2)
    add_para(doc, "当前项目使用访问令牌，没有实现 refresh token、令牌黑名单和密钥轮换；application.yml 中提供了开发环境默认密钥。生产环境应通过安全的环境变量注入高强度密钥，并增加刷新令牌、撤销机制和更细粒度的审计日志。答辩时主动说明这些边界，通常比声称系统“绝对安全”更专业。")


def section_appointment(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "5. 服务预约、时间段与技师推荐", 1)
    add_heading(doc, "5.1 预约为什么是项目核心", 2)
    add_para(doc, "预约模块连接用户、服务项目、技师、排班、请假、支付、通知、聊天和评价，是系统业务耦合最集中的部分。真正的难点不是插入一条预约记录，而是保证技师具有服务能力、完整处于排班时间内、没有请假、没有重叠预约，并在并发请求下尽量避免重复占用。")

    add_heading(doc, "5.2 创建预约的关键代码", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/AppointmentService.java:169",
        """
ServiceItem item = serviceItemOrThrow(req.getServiceItemId());
LocalDateTime endTime =
        req.getAppointmentTime().plusMinutes(item.getDurationMin());

Long technicianId = req.getTechnicianId();
if (technicianId == null) {
    technicianId = schedulingService
            .recommendTechnician(item.getId(), req.getAppointmentTime())
            .getTechnicianId();
}

String lockKey = "lock:appointment:" + technicianId + ":" + date;
String token = redisLockService.tryLock(lockKey, Duration.ofSeconds(8));
try {
    scheduleConflictChecker.validateTechnicianForAppointment(
            technicianId, item.getId(), startTime, endTime);
    if (scheduleConflictChecker.hasConflictInRange(
            technicianId, startTime, endTime)) {
        throw new BizException(4013, "该时间段技师已被预约");
    }
    appointment.setStatus("CREATED");
    appointment.setPayStatus("UNPAID");
    appointmentMapper.insert(appointment);
} finally {
    redisLockService.unlock(lockKey, token);
}
""",
        "服务时长参与结束时间计算；未指定技师时自动推荐；先加锁再二次校验；预约初始状态为 CREATED，支付状态为 UNPAID。",
    )

    add_heading(doc, "5.3 为什么要保存服务时长快照", 2)
    add_para(doc, "预约表保存 appointment_end_time 和 duration_min_snapshot。假设用户预约时服务时长为 60 分钟，管理员后来把服务改成 90 分钟，历史预约不应跟着变化。快照保证历史订单能够按创建时的业务信息解释，同时兼容旧数据时仍可回退读取当前服务时长或默认 60 分钟。")

    add_heading(doc, "5.4 时间区间重叠算法", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/mapper/AppointmentMapper.java:16",
        """
WHERE technician_id = #{technicianId}
  AND appointment_time < #{endTime}
  AND COALESCE(appointment_end_time,
      DATE_ADD(appointment_time,
        INTERVAL COALESCE(duration_min_snapshot, 60) MINUTE)) > #{startTime}
  AND status IN ('PENDING', 'APPROVED', 'WAIT_SERVICE', 'IN_SERVICE')
""",
        "两个半开区间 [A.start, A.end) 与 [B.start, B.end) 重叠的条件是 A.start < B.end 且 A.end > B.start。相邻预约前一个结束等于后一个开始时不会被判为冲突。",
    )

    add_heading(doc, "5.5 Redis 锁的实现与边界", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/cache/RedisLockService.java",
        """
String token = UUID.randomUUID().toString();
Boolean ok = redisTemplate.opsForValue()
        .setIfAbsent(key, token, ttl);

// 解锁时只有 value 仍等于自己的 token 才删除
if redis.call('get', KEYS[1]) == ARGV[1] then
    return redis.call('del', KEYS[1])
end
""",
        "setIfAbsent 对应 Redis SET NX，并设置过期时间避免死锁；Lua 解锁保证“比较 token 与删除 key”原子执行，避免误删别人的锁。",
    )
    add_callout(doc, "设计取舍：", "当前锁粒度是“技师 + 日期”，实现简单且能覆盖同一天的冲突检查，但会让同一技师当天不同时间段的创建请求串行。可进一步按更细时间片加锁，或配合数据库唯一约束与版本控制。", "warn")

    add_heading(doc, "5.6 技师推荐算法", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/SchedulingService.java",
        """
return candidates.stream()
    .filter(t -> isWorkingDuring(t, start, end))
    .filter(t -> !isOnLeaveDuring(t, start, end))
    .filter(t -> !hasConflictInRange(t, start, end))
    .map(t -> toRecommendation(t, start, end, duration))
    .max(comparingDouble(TechnicianRecommendationVo::getScore))
    .orElse(null);

double score = 60 + rating * 8 - dayLoad * 10;
""",
        "先做硬约束过滤，再做软评分排序。评分越高越优先，当日订单越多则扣分，从而兼顾服务质量与工作量均衡。",
    )

    add_heading(doc, "5.7 可预约时间段生成", 2)
    add_para(doc, "系统以 30 分钟为步长遍历技师排班。最晚开始时间等于排班结束时间减去服务时长；每个候选时间段都检查当前时间、请假和预约冲突，并汇总该时段的可用技师数量与推荐技师。")

    add_heading(doc, "5.8 预约状态机", 2)
    add_table(
        doc,
        ["状态", "业务含义", "允许的主要下一步"],
        [
            ("CREATED", "预约已创建，尚未支付", "支付成功 -> PENDING；用户取消 -> CANCELLED"),
            ("PENDING", "已支付，等待技师或管理员确认", "APPROVED 或 REJECTED"),
            ("APPROVED", "预约已确认", "WAIT_SERVICE"),
            ("WAIT_SERVICE", "等待到店开始服务", "IN_SERVICE"),
            ("IN_SERVICE", "服务进行中", "COMPLETED"),
            ("COMPLETED", "服务完成，可评价", "终态"),
            ("REJECTED / CANCELLED", "预约被拒绝或取消", "必要时触发退款，终态"),
        ],
        [1600, 4100, 3660],
        center_cols={0},
    )
    add_callout(doc, "为什么要状态机：", "状态机把业务步骤写成白名单。即使接口被重复或乱序调用，也不能从 PENDING 直接跳到 COMPLETED。", "good")


def section_schedule(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "6. 排班、请假与考勤协同", 1)
    add_heading(doc, "6.1 三者的关系", 2)
    add_para(doc, "排班描述技师计划工作时间，请假描述临时不可用区间，考勤记录实际到岗情况。预约可用性主要依赖排班与已批准请假；考勤用于技师工作管理和统计。三者不能各自独立，否则会出现管理员批准请假后仍有用户预约、或覆盖排班时删除已有预约承载时间的问题。")

    add_heading(doc, "6.2 请假审批的冲突判断", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/LeaveRecordService.java:269",
        """
if ("APPROVED".equals(targetStatus)) {
    ensureNoApprovedLeaveConflict(
        technicianId, startTime, endTime, record.getId());
    ensureNoAppointmentConflict(
        technicianId, startTime, endTime);
}
record.setStatus(targetStatus);
record.setReviewedByUserId(currentUser.getUserId());
record.setReviewedAt(LocalDateTime.now());
leaveRecordMapper.updateById(record);
notificationService.notifyUser(...);
""",
        "管理员批准前同时检查其他已批准请假和有效预约。审批结果记录审核人、审核时间与备注，并通知技师。",
    )

    add_heading(doc, "6.3 冲突预览与改派候选", 2)
    add_para(doc, "管理员可以先预览请假时间内冲突的预约。系统会为每条冲突预约搜索同一服务能力、同一时间有排班、没有请假且没有其他预约冲突的技师，形成候选改派列表。这说明系统不仅发现冲突，还提供解决冲突所需的数据。")

    add_heading(doc, "6.4 自动排班算法", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/AutoScheduleService.java:106",
        """
List<Technician> candidates = technicians.stream()
    .filter(t -> canUseTechnician(req, date, t.getId()))
    .sorted(comparingLong(t ->
            historicalLoad(t) + generatedLoad(t))
        .thenComparingInt(t -> rotationRank(t, date))
        .thenComparing(t -> -rating(t))
        .thenComparing(Technician::getId))
    .limit(req.getDailyTechnicianCount())
    .toList();
""",
        "先排除请假、已有排班和预约冲突，再按近期历史负载、当前生成负载、周期轮换、评分和 ID 排序，优先选择工作量较少的技师。",
    )
    add_callout(doc, "为什么先预览再应用：", "preview 让管理员看到生成数量、缺口和警告；apply 复用预览结果写入数据库。这个两阶段设计降低批量排班直接落库的风险。", "info")

    add_heading(doc, "6.5 考勤模块答辩表达", 2)
    add_para(doc, "考勤记录以“技师 + 日期”唯一，保存签到、签退和状态。技师端工作台可展示当日状态和月度统计。它与预约模块的主要关系是运营分析，而不是直接决定用户是否可以预约；预约可用性仍以排班、请假和冲突为核心。")


def section_payment(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "7. 统一支付、钱包与退款", 1)
    add_heading(doc, "7.1 为什么需要统一支付单", 2)
    add_para(doc, "项目中有三类需要支付的业务：商品订单、服务预约和钱包充值。如果各模块分别实现支付宝、微信、超时关闭和退款，会产生大量重复逻辑。PaymentService 把它们统一抽象为 PaymentOrder，通过 business_type 和 business_id 指向具体业务。")
    add_table(
        doc,
        ["business_type", "业务对象", "支付成功后的回写"],
        [
            ("PRODUCT_ORDER", "商品订单", "pay_status=PAID，status=PAID，通知用户与管理员"),
            ("APPOINTMENT", "服务预约", "pay_status=PAID，正常时 status=PENDING，通知用户与技师"),
            ("WALLET_RECHARGE", "钱包充值", "增加余额，记录充值与赠送流水"),
        ],
        [2200, 2600, 4560],
        center_cols={0},
    )

    add_heading(doc, "7.2 创建支付单", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/PaymentService.java:135",
        """
closePendingPayments("APPOINTMENT", appointment.getId(),
        "A new payment was created");

PaymentOrder paymentOrder = new PaymentOrder();
paymentOrder.setBusinessType("APPOINTMENT");
paymentOrder.setBusinessId(appointment.getId());
paymentOrder.setPayChannel(payChannel);
paymentOrder.setPayMode(payMode);
paymentOrder.setStatus("PAYING");
paymentOrder.setExpireTime(LocalDateTime.now().plusMinutes(15));
paymentOrderMapper.insert(paymentOrder);

return gatewayForMode(payMode).create(paymentOrder);
""",
        "重新发起支付会先关闭旧的待支付单；支付单有效期为 15 分钟；网关创建结果可能是二维码、支付宝跳转表单或手动确认信息。",
    )

    add_heading(doc, "7.3 网关路由", 2)
    add_para(doc, "PaymentGateway 定义统一接口，PaymentGatewayRouter 根据 pay_mode 解析具体实现。项目提供 AlipayPagePaymentGateway、WechatNativePaymentGateway 和 ManualPaymentGateway。开发或答辩环境可使用 MANUAL_CASHIER 演示完整业务，不依赖真实商户资质；配置齐全时可切换外部网关。")
    add_callout(doc, "答辩措辞：", "不要把手动收银说成“假支付”。更准确的说法是：系统通过网关接口隔离支付渠道，毕业设计演示环境默认使用手动收银实现可重复测试，真实支付宝和微信网关已有对应适配类与配置入口。", "warn")

    add_heading(doc, "7.4 支付回调为什么要幂等", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/PaymentService.java:457",
        """
if (!PENDING_STATUSES.contains(paymentOrder.getStatus())) {
    return;
}
int updated = paymentOrderMapper.markSuccessIfPending(
        paymentOrder.getId(), platformTradeNo, paidTime, notifyContent);
if (updated == 0) {
    return;
}
applyBusinessSuccess(paymentOrderMapper.selectById(paymentOrder.getId()));
""",
        "支付平台可能重复通知。只有支付单仍处于待处理状态时才更新为成功，受影响行数为 0 就直接返回，避免重复增加钱包余额或重复推进订单。",
    )

    add_heading(doc, "7.5 预约支付后的二次档期校验", 2)
    add_para(doc, "外部支付存在时间差：用户创建预约后可能过一段时间才完成支付。支付成功时，系统通过 appointmentPaymentBlockReason 再次确认技师档期。如果此时档期已不可用，预约会被取消并发起退款，同时通知用户。这个设计避免“钱已付但服务无法履约”。")
    add_callout(doc, "这是强答辩点：", "创建预约时校验一次，支付成功时再校验一次，分别处理提交瞬间的并发和外部支付延迟造成的状态变化。", "good")

    add_heading(doc, "7.6 超时关闭与退款", 2)
    add_para(doc, "PaymentReconciliationJob 定期调用 closeExpiredPayments。系统先查询远端状态，若远端已成功则补做成功回写；否则尝试关闭网关支付单，再关闭本地支付单并让业务过期。退款通过 PaymentRefund 记录，钱包余额支付或手动收银可退回钱包，外部渠道则调用相应网关或进入人工处理状态。")

    add_heading(doc, "7.7 钱包流水", 2)
    add_para(doc, "UserWallet 保存余额与累计充值、赠送金额，WalletTxn 保存每次充值、消费和退款后的余额。txn_no 具有唯一性，可防止重复记账。余额支付商品或预约时由 WalletService 扣款，退款时按业务类型记录对应流水。")


def section_mall(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "8. 商品商城、库存与售后", 1)
    add_heading(doc, "8.1 商城闭环", 2)
    add_code(
        doc,
        "商品订单状态",
        """
浏览商品 -> 加入购物车/直接购买 -> 选择地址 -> 创建订单
CREATED(待支付) -> PAID(待发货) -> SHIPPED(待收货)
-> COMPLETED(已完成)
异常分支：CANCELLED / REFUNDED
""",
        "订单主表保存金额、支付、收货地址快照、物流和售后状态；订单明细保存商品、数量与成交价快照。",
    )

    add_heading(doc, "8.2 为什么订单要保存地址快照", 2)
    add_para(doc, "用户下单后可能修改或删除常用地址。如果订单只保存 address_id，历史订单展示会变化甚至无法查询。因此创建订单时复制收件人、电话、省市区、详细地址和标签，保证订单历史稳定。")

    add_heading(doc, "8.3 库存锁与扣减", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/ProductOrderService.java:139",
        """
for (OrderLine line : lines) {
    String key = "lock:product:stock:" + line.productId();
    String token = redisLockService.tryLock(key, Duration.ofSeconds(8));
    lockTokens.put(key, token);

    Product latest = productMapper.selectById(line.productId());
    if (latest.getStock() < line.quantity()) {
        throw new BizException(4102, "商品库存不足");
    }
    latest.setStock(latest.getStock() - line.quantity());
    productMapper.updateById(latest);
}
""",
        "锁按商品 ID 划分，加锁后重新读取最新库存再扣减，避免两个请求都基于旧库存判断。事务失败时数据库修改回滚，finally 释放全部锁。",
    )
    add_callout(doc, "可以进一步改进：", "当前扣库存采用“查询后更新”。更强的数据库方案是使用 UPDATE product SET stock=stock-? WHERE id=? AND stock>=?，根据影响行数判断是否成功，再与 Redis 锁形成双重保护。", "warn")

    add_heading(doc, "8.4 取消订单为什么要恢复库存", 2)
    add_para(doc, "订单创建时已经占用库存。用户在未支付状态取消、支付单过期导致业务关闭或售后退款通过时，需要调用 restoreStock 将订单明细数量加回库存，否则库存会永久减少。库存恢复必须与订单状态变化放在同一事务中。")

    add_heading(doc, "8.5 售后审核", 2)
    add_para(doc, "用户只可对已支付未发货或已完成订单申请退款/售后。管理员审核通过后恢复库存、更新售后状态和订单支付状态、创建退款记录并通知用户；审核拒绝则记录处理备注。商品订单状态与退款状态分开保存，便于展示“订单业务进度”和“资金处理进度”。")


def section_other_modules(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "9. AI 顾问、聊天、通知与统计", 1)
    add_heading(doc, "9.1 AI 顾问不是孤立聊天", 2)
    add_para(doc, "AiConsultService 会读取启用中的服务和商品目录，把用户文字、可选图片、历史咨询上下文和真实目录一起发送给兼容 OpenAI 协议的模型。模型返回的服务 ID 和商品 ID 必须再次映射到本地目录，过滤不存在、重复或不可用项。")
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/AiConsultService.java:230",
        """
Intent intent = detectIntent(question);
AiBeautyAnalysis remote =
        tryRemoteAnalysis(question, imageUrl, record, catalog);

List<ServiceRecommendation> services =
        mergeServiceRecommendations(remote, catalog.servicesById());
List<ProductRecommendation> products =
        mergeProductRecommendations(remote, catalog.productsById());

if (intent.serviceIntent() && services.isEmpty()) {
    services = recommendServices(question, catalog.services());
}
if (intent.productIntent() && products.isEmpty()) {
    products = recommendProducts(question, catalog.products());
}
String mode = remote == null ? "LOCAL_RULE" : "REMOTE_API";
""",
        "远程模型负责理解与生成，本地目录负责事实约束；本地规则可按关键词评分推荐，并记录 REMOTE_API 或 LOCAL_RULE 便于排查。",
    )
    add_callout(doc, "AI 风险控制：", "系统会显示建议仅供参考，实际方案需结合到店沟通、肤质发质和过敏情况。生产环境还应加强图片隐私、内容安全、调用限流和用户授权。", "risk")

    add_heading(doc, "9.2 聊天消息的落库与推送", 2)
    add_code(
        doc,
        "backend/src/main/java/com/bfsm/service/ChatService.java:204",
        """
ChatSession session = findOrCreateSession(senderId, receiverId);
ChatMessage message = new ChatMessage();
message.setSessionId(session.getId());
message.setSenderId(senderId);
message.setReceiverId(receiverId);
message.setContent(content.trim());
chatMessageMapper.insert(message);
chatWebSocketHandler.pushTo(receiverId, message);
""",
        "消息先写数据库，再通过 WebSocket 推送。即使接收者不在线，之后仍可从历史消息中查询。",
    )

    add_heading(doc, "9.3 通知与聊天的区别", 2)
    add_table(
        doc,
        ["机制", "适合内容", "项目示例"],
        [
            ("系统通知", "业务状态提醒，强调可追踪与未读状态", "预约支付成功、请假审核、订单发货、售后处理"),
            ("即时聊天", "用户与技师之间的双向沟通", "预约确认后系统消息、用户咨询服务细节"),
        ],
        [1800, 3800, 3760],
    )

    add_heading(doc, "9.4 统计模块", 2)
    add_para(doc, "StatsService 将预约、订单、服务、商品和技师数据聚合为管理看板与技师工作台数据，包括预约数量、销售额、服务统计、商品排行、销售趋势、技师业绩和考勤图表。前端使用 ECharts 展示趋势和分类。统计模块的价值是把业务操作沉淀为运营决策信息。")


def section_database(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "10. 数据库设计与实体关系", 1)
    add_heading(doc, "10.1 数据表分组", 2)
    add_table(
        doc,
        ["领域", "核心表", "关系说明"],
        [
            ("用户权限", "sys_user、sys_role、sys_user_role、technician", "用户与角色多对多；技师扩展用户业务属性"),
            ("服务与预约", "service_category、service_item、technician_service、schedule、appointment、appointment_log", "服务分类与项目；技师服务能力；排班与预约；日志记录状态变化"),
            ("请假考勤", "leave_record、attendance_record", "请假区间参与预约可用性；考勤按技师和日期唯一"),
            ("商城订单", "product_category、product、shopping_cart_item、user_address、product_order、product_order_item", "订单主从结构，并保存地址和价格快照"),
            ("支付钱包", "user_wallet、wallet_txn、recharge_activity、payment_order、payment_refund", "钱包账户与流水；统一支付单和退款单"),
            ("互动消息", "favorite、like_record、comment、system_notification、chat_session、chat_message", "收藏点赞评价、业务通知、即时会话"),
            ("AI 咨询", "ai_consult_record、ai_consult_message 等", "由 AI 初始化器维护咨询会话、消息和推荐快照"),
        ],
        [1450, 4650, 3260],
    )

    add_heading(doc, "10.2 最重要的关系链", 2)
    add_code(
        doc,
        "核心关系示意",
        """
sys_user --< sys_user_role >-- sys_role
sys_user --1 technician
technician >--< service_item
technician --< schedule
sys_user --< appointment >-- technician
appointment >-- service_item
appointment --< appointment_log

sys_user --< product_order --< product_order_item >-- product
payment_order --(business_type + business_id)--> appointment/product_order/wallet
payment_order --< payment_refund
""",
        "预约是用户、技师和服务项目的关联实体；商品订单使用主从表；统一支付通过业务类型和业务 ID 建立逻辑关联。",
    )

    add_heading(doc, "10.3 为什么同时使用外键、状态和日志", 2)
    add_para(doc, "外键保证关联对象存在；状态字段表达对象当前业务阶段；日志记录对象如何到达当前阶段。以预约为例，appointment.status 告诉页面当前状态，appointment_log 保存每次 from_status、to_status、操作人和备注，便于展示进度与追溯。")

    add_heading(doc, "10.4 索引与性能思考", 2)
    add_para(doc, "当前表已经对用户名、手机号、支付单号和部分唯一关系建立约束。随着数据量增加，预约冲突查询常用 technician_id、appointment_time、status，通知查询常用 user_id、read_status、created_at，订单查询常用 user_id、status、created_at，这些组合应建立合适索引。答辩时可以说明项目当前以毕业设计规模为主，生产化需要结合慢查询分析补充组合索引。")


def section_debugging(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "11. 代码阅读、调试与风险认识", 1)
    add_heading(doc, "11.1 推荐的源码阅读顺序", 2)
    for text in [
        "从 frontend/src/router/index.ts 看三类角色有哪些页面。",
        "从某个页面找到调用的 API 方法，再定位后端 Controller 的 URL。",
        "阅读 Controller 的请求 DTO、角色注解和调用的 Service 方法。",
        "在 Service 中识别事务边界、状态判断、权限判断和外部依赖。",
        "查看 Mapper 的 SQL 与对应 Entity，最后回到 schema.sql 理解字段。",
        "遇到支付、Redis、AI 或 WebSocket，再进入 gateway、cache、service/ai、websocket 子目录。",
    ]:
        add_number(doc, text)

    add_heading(doc, "11.2 常用运行与检查命令", 2)
    add_code(
        doc,
        "本地开发",
        r"""
# 后端
$env:JAVA_HOME='D:\IDEA\JAVA\JDK17'
mvn -f backend/pom.xml test
mvn -f backend/pom.xml spring-boot:run

# 前端
Set-Location frontend
npm install
npm run dev
npm run build

# 接口文档
http://localhost:8080/doc.html
""",
        "后端依赖 MySQL 和 Redis；前端默认请求 http://localhost:8080。支付默认可使用 MANUAL_CASHIER，AI 远程调用需要配置 API Key。",
    )

    add_heading(doc, "11.3 一条错误如何定位", 2)
    add_table(
        doc,
        ["现象", "优先检查", "典型原因"],
        [
            ("页面提示 401", "localStorage token、Axios 请求头、后端 JWT 密钥", "令牌过期、后端重启后密钥变化、请求未携带 Bearer"),
            ("预约没有可用时段", "schedule、leave_record、appointment、service duration", "未排班、已请假、时长超出排班、存在区间冲突"),
            ("预约支付后退款", "PaymentService 日志与 appointmentPaymentBlockReason", "支付期间技师档期发生变化"),
            ("商品库存不足", "product.stock 与库存锁", "并发订单占用库存或测试数据不足"),
            ("AI 请求失败", "application.yml、API Key、降级开关", "远程配置缺失、网络问题、禁止降级"),
            ("聊天不实时", "WebSocket 握手 token、在线连接、消息是否已落库", "连接未建立，但历史消息通常仍可查询"),
        ],
        [2100, 3700, 3560],
    )

    add_heading(doc, "11.4 应主动说明的工程改进方向", 2)
    for text in [
        "敏感配置：数据库密码、JWT 默认密钥和支付密钥应全部使用环境变量或密钥管理服务。",
        "并发一致性：库存可增加数据库条件更新；预约可增加更细锁粒度或数据库层约束。",
        "可观测性：增加统一 traceId、结构化日志、支付回调审计和关键业务指标。",
        "测试覆盖：现有 Service 测试覆盖核心模块，可进一步增加 Controller 集成测试、并发测试和支付回调幂等测试。",
        "前端安全：避免长期把高价值令牌放在可被脚本读取的位置，生产环境可评估 HttpOnly Cookie 与 CSRF 防护组合。",
        "AI 合规：增加图片授权、内容过滤、数据保留策略和模型输出评估。",
    ]:
        add_bullet(doc, text)


QA = [
    ("为什么选择 Spring Boot？", "它提供成熟的 Web、依赖注入、事务、安全和配置能力，能让我把精力集中在预约冲突、支付回写和状态流转等业务问题上。项目使用 Spring Boot 3.4.4 与 Java 17。"),
    ("为什么不用 Session，而用 JWT？", "前后端分离下，JWT 让每次请求携带身份信息，后端不依赖服务器会话，部署和接口调用更简单。代价是令牌撤销和续期需要额外设计，当前项目尚未实现 refresh token。"),
    ("前端已经限制角色页面，为什么后端还要鉴权？", "前端代码运行在用户设备上，可以被绕过。后端的 SecurityConfig、@PreAuthorize 和数据归属判断才是真正安全边界。"),
    ("如何防止同一技师同一时间被重复预约？", "系统先检查服务能力、排班、请假和预约区间冲突，再以技师和日期为 Redis 锁键获取锁，锁内重新校验后才写入预约。"),
    ("预约冲突算法是什么？", "两个时间区间重叠的条件是旧预约开始时间小于新预约结束时间，并且旧预约结束时间大于新预约开始时间。查询只统计待确认、已确认、待服务和服务中等有效状态。"),
    ("为什么需要服务时长快照？", "服务时长可能被管理员修改，历史预约必须保留创建时的结束时间和时长，否则历史记录和冲突判断会被后续配置改变。"),
    ("技师推荐是否用了机器学习？", "没有。当前采用可解释的规则评分：先过滤不可用技师，再按评分和当天负载计算分数。对毕业设计场景来说可解释、易验证，未来可基于历史数据训练推荐模型。"),
    ("为什么创建预约后不是直接 PENDING？", "CREATED 表示待支付，支付成功后才进入 PENDING 等待确认。这样能区分占位但未完成支付的请求，支付超时也能关闭业务。"),
    ("支付成功为什么还要检查档期？", "外部支付存在延迟，支付期间档期可能变化。成功回调时再次检查，若不可履约就取消预约并发起退款，避免用户付费后无法服务。"),
    ("如何防止支付回调重复处理？", "只有待处理支付单可以更新成功，Mapper 使用条件更新并检查影响行数。已经成功的支付单再次收到通知会直接返回。"),
    ("手动收银模式有什么意义？", "它是 PaymentGateway 的一个实现，用于没有真实商户配置时稳定演示支付闭环。业务模块不依赖具体渠道，配置完整后可切换支付宝或微信网关。"),
    ("如何防止商品超卖？", "按商品 ID 获取 Redis 库存锁，加锁后重新读取库存并扣减；事务异常会回滚，finally 释放锁。进一步可用数据库条件更新增强一致性。"),
    ("为什么订单要保存地址和价格快照？", "用户地址、商品价格会变化，订单必须保留成交当时的信息，保证历史可追溯。"),
    ("请假与预约冲突如何处理？", "管理员审批请假前检查有效预约；冲突预览会列出预约并计算可改派技师。只有解决冲突后才适合批准请假。"),
    ("自动排班的公平性体现在哪里？", "系统优先选择历史和本轮生成负载较少的技师，并加入周期轮换，再参考评分，避免长期固定选择同一批人。"),
    ("AI 推荐如何避免推荐不存在的服务？", "远程模型返回的 ID 会与数据库当前启用目录匹配，无法匹配或重复的结果被过滤；没有结果时按关键词和目录做本地规则推荐。"),
    ("聊天为什么要先落库再推送？", "WebSocket 只保证在线实时性。先落库可以保证离线用户之后仍能看到消息，也便于会话列表和历史查询。"),
    ("事务用在哪里？", "预约创建、订单创建、支付回写、请假审批、钱包扣款和退款等多表操作使用 @Transactional，确保一组数据库修改要么全部成功，要么回滚。"),
    ("项目最大的不足是什么？", "我认为主要是生产化能力还可加强，包括 refresh token、密钥管理、组合索引、并发压力测试、支付审计和 AI 隐私治理。核心业务链路已实现，但仍有工程化提升空间。"),
    ("你个人完成了什么？", "我围绕需求分析、数据库设计、前后端实现和测试完成了多角色业务，并重点实现预约冲突控制、统一支付、库存并发、排班请假协同和 AI 目录推荐等核心模块。回答时应结合自己实际完成情况补充。"),
]


def section_qa(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "12. 高频答辩问题与参考回答", 1)
    add_para(doc, "下面的回答采用“结论 + 实现 + 取舍”结构。现场不必逐字背诵，但应保留关键类名、状态值和设计原因。")
    for idx, (question, answer) in enumerate(QA, 1):
        add_heading(doc, f"12.{idx} {question}", 2)
        add_para(doc, answer)


def section_demo(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "13. 现场演示脚本与故障预案", 1)
    add_heading(doc, "13.1 推荐的 10 分钟演示顺序", 2)
    demo_steps = [
        ("0:00-0:45", "登录与角色", "展示普通用户登录，说明 JWT；快速指出技师端和管理端入口不同。"),
        ("0:45-2:30", "服务预约", "选择服务、日期和时间，展示可用时段与推荐技师，创建预约。"),
        ("2:30-3:30", "支付", "打开收银台，使用手动收银或当前可用渠道完成支付，说明统一支付单。"),
        ("3:30-4:30", "技师履约", "切换技师账号，确认预约并推进到待服务/服务中。"),
        ("4:30-5:30", "用户闭环", "查看预约详情、状态日志和通知，完成后展示评价入口。"),
        ("5:30-7:00", "商城订单", "商品、购物车、地址、下单、支付；管理端展示发货与售后。"),
        ("7:00-8:00", "排班请假", "技师提交请假，管理员查看冲突或审批，说明与预约联动。"),
        ("8:00-9:00", "AI 与聊天", "展示 AI 推荐如何关联真实服务/商品，再展示聊天或通知。"),
        ("9:00-10:00", "统计与总结", "管理看板展示预约、销售和技师统计，回到四个技术亮点收尾。"),
    ]
    add_table(doc, ["时间", "演示内容", "讲解重点"], demo_steps, [1400, 2300, 5660], center_cols={0, 1})

    add_heading(doc, "13.2 演示前检查清单", 2)
    for item in [
        "MySQL 的 bfsm 数据库可连接，schema.sql 已初始化。",
        "Redis 已启动，否则缓存和分布式锁相关功能可能失败。",
        "后端 8080 端口和前端 5173 端口可访问。",
        "准备 USER、TECH、ADMIN 三类账号，并提前确认密码。",
        "为演示日期准备技师排班，避免预约页面没有时段。",
        "商品库存大于 0，用户有收货地址，手动收银开关可用。",
        "AI 未配置时提前决定是否启用本地降级，避免现场等待远程接口。",
        "浏览器预先打开用户端、技师端、管理端三个标签页。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "13.3 常见故障预案", 2)
    add_table(
        doc,
        ["故障", "现场处理", "解释方式"],
        [
            ("预约没有时间段", "切换到预先安排排班的日期或技师", "可预约时段依赖服务时长、排班、请假和冲突，不是固定写死"),
            ("支付外部渠道不可用", "切换 MANUAL_CASHIER 完成流程", "网关抽象允许不同环境替换渠道"),
            ("AI 远程超时", "使用本地规则降级或展示已有咨询记录", "AI 能力受外部服务影响，系统设计了降级路径"),
            ("WebSocket 未连接", "刷新聊天页面并展示消息历史", "消息先落库，实时推送失败不影响历史持久化"),
            ("统计图数据少", "使用初始化数据或先完成一笔预约/订单", "统计来自真实业务数据聚合"),
        ],
        [2100, 3200, 4060],
    )

    add_heading(doc, "13.4 结束陈述", 2)
    add_callout(
        doc,
        "收尾可直接说：",
        "本项目围绕美容美发门店的实际经营流程，形成了预约、排班、履约、商城、支付、消息和统计等业务闭环。技术上采用前后端分离架构，使用 JWT 实现多角色认证，MyBatis-Plus 操作 MySQL，Redis 处理缓存和并发锁，WebSocket 提供实时通信。实现过程中我重点解决了时间区间冲突、状态合法流转、库存并发扣减、统一支付回写和外部 AI 结果约束等问题。",
        "good",
    )


def section_cheatsheet(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "14. 最后一页速记", 1)
    add_table(
        doc,
        ["记忆点", "必须说出的关键词"],
        [
            ("架构", "Vue 3 + TypeScript + Spring Boot 3.4.4 + Java 17 + MySQL + Redis"),
            ("认证", "BCrypt 密码哈希；JWT 无状态；JwtAuthFilter；@PreAuthorize"),
            ("预约", "服务时长快照；排班/请假/冲突；Redis 技师日期锁；锁内二次校验"),
            ("冲突公式", "old.start < new.end 且 old.end > new.start"),
            ("推荐", "硬约束过滤；score = 60 + rating*8 - dayLoad*10"),
            ("预约状态", "CREATED -> PENDING -> APPROVED -> WAIT_SERVICE -> IN_SERVICE -> COMPLETED"),
            ("商品状态", "CREATED -> PAID -> SHIPPED -> COMPLETED；异常 CANCELLED/REFUNDED"),
            ("支付", "PaymentOrder 统一商品、预约、充值；网关路由；15 分钟过期；幂等回调"),
            ("库存", "商品 ID 锁；锁后重读；事务回滚；取消/退款恢复库存"),
            ("请假排班", "审批前查预约；冲突预览；候选改派；预览后应用自动排班"),
            ("AI", "远程模型 + 真实目录约束 + 本地规则降级 + 风险提示"),
            ("消息", "通知负责业务状态；聊天先落库后 WebSocket 推送"),
            ("数据库", "31 张主脚本业务表；预约日志；订单主从；支付与退款分离"),
            ("不足", "refresh token、密钥管理、索引、并发压测、可观测性、AI 隐私"),
        ],
        [1900, 7460],
    )
    add_callout(doc, "答题公式：", "业务问题是什么 -> 系统采用什么方案 -> 关键类或表在哪里 -> 这样设计的优点与边界是什么。", "warn")
    add_para(doc, "建议最后再打开源码，对照以下五个入口复习：SecurityConfig、AppointmentService、PaymentService、ProductOrderService、AiConsultService。能够讲清这五个类，基本就能覆盖项目答辩的大部分技术追问。")


def add_appendix(doc: Document) -> None:
    page_break(doc)
    add_heading(doc, "附录：核心源码定位表", 1)
    rows = [
        ("登录注册", "AuthController / AuthService / JwtTokenService / JwtAuthFilter"),
        ("权限配置", "SecurityConfig / frontend/src/router/index.ts"),
        ("预约入口", "AppointmentController / AppointmentService"),
        ("技师可用性", "ScheduleConflictChecker / SchedulingService"),
        ("自动排班", "AutoScheduleService / SchedulingController"),
        ("请假考勤", "LeaveRecordService / AttendanceService"),
        ("商品订单", "ProductOrderService / ProductController"),
        ("支付退款", "PaymentService / PaymentController / service/gateway"),
        ("钱包", "WalletService / WalletController"),
        ("AI 咨询", "AiConsultService / service/ai / AiConsultController"),
        ("聊天通知", "ChatService / ChatWebSocketHandler / NotificationService"),
        ("后台统计", "StatsService / StatsController"),
        ("数据库", "backend/src/main/resources/schema.sql"),
        ("前端请求", "frontend/src/utils/request.ts / api/modules.ts"),
        ("前端角色页面", "frontend/src/views/user、tech、admin"),
    ]
    add_table(doc, ["模块", "主要源码位置"], rows, [2200, 7160])
    add_para(doc, "本文档依据 2026 年 6 月 9 日工作区中的当前源码与论文初稿整理。后续若代码继续修改，应以实际实现为准，尤其注意状态值、接口地址和支付配置是否发生变化。", color=MUTED, italic=True)


def build() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)
    section_project_overview(doc)
    section_architecture(doc)
    section_request_flow(doc)
    section_security(doc)
    section_appointment(doc)
    section_schedule(doc)
    section_payment(doc)
    section_mall(doc)
    section_other_modules(doc)
    section_database(doc)
    section_debugging(doc)
    section_qa(doc)
    section_demo(doc)
    section_cheatsheet(doc)
    add_appendix(doc)

    # Keep major chapters visually separated and avoid orphan headings.
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
