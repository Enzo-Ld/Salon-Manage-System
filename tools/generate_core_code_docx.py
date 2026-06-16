from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "docs" / "appendix_core_source_code.md"
OUTPUT_DOCX = ROOT / "docs" / "BFSM核心代码实现说明.docx"


def set_run_font(run, name: str, east_asia: str | None = None, size: Pt | None = None):
    run.font.name = name
    if size is not None:
        run.font.size = size
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if east_asia:
        rfonts.set(qn("w:eastAsia"), east_asia)


def set_paragraph_shading(paragraph, fill: str):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_bottom_border(paragraph, color: str = "D9E2F3"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(32, 32, 32)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("CodeBlockLine", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor(40, 40, 40)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.04)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    source = styles.add_style("SourcePath", 1)
    source.font.name = "Calibri"
    source.font.size = Pt(9)
    source.font.italic = True
    source.font.color.rgb = RGBColor(90, 90, 90)
    source._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    source.paragraph_format.space_before = Pt(2)
    source.paragraph_format.space_after = Pt(6)

    note = styles.add_style("GuideNote", 1)
    note.font.name = "Calibri"
    note.font.size = Pt(10)
    note.font.color.rgb = RGBColor(60, 60, 60)
    note._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    note.paragraph_format.left_indent = Inches(0.12)
    note.paragraph_format.right_indent = Inches(0.12)
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(8)
    note.paragraph_format.line_spacing = 1.25


def add_footer(doc: Document):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BFSM 核心代码实现说明")
    set_run_font(run, "Calibri", "Microsoft YaHei", Pt(9))
    run.font.color.rgb = RGBColor(120, 120, 120)


def add_cover(doc: Document):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("BFSM 美容美发服务管理系统")
    set_run_font(run, "Calibri", "Microsoft YaHei", Pt(20))
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("1F4D78")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("核心代码实现说明")
    set_run_font(run, "Calibri", "Microsoft YaHei", Pt(16))
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("2E74B5")

    note = doc.add_paragraph(style="GuideNote")
    set_paragraph_shading(note, "F4F6F9")
    run = note.add_run(
        "文档说明：本文件根据项目实际源码整理，覆盖认证权限、预约排班、支付、商品订单、Redis 并发锁、"
        "AI 咨询以及前端请求与路由控制等核心模块。代码块中以“附录注释”开头的内容为说明性注释，"
        "用于解释关键业务规则、并发控制和状态流转逻辑。"
    )
    set_run_font(run, "Calibri", "Microsoft YaHei", Pt(10))

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("主要模块：").bold = True
    for item in [
        "认证与权限模块",
        "预约与排班模块",
        "支付模块",
        "商品订单模块",
        "Redis 分布式锁",
        "AI 咨询模块",
        "前端请求与路由控制",
    ]:
        bullet = doc.add_paragraph(style="List Bullet")
        bullet.paragraph_format.left_indent = Inches(0.38)
        run = bullet.add_run(item)
        set_run_font(run, "Calibri", "Microsoft YaHei", Pt(11))

    doc.add_section(WD_SECTION.NEW_PAGE)


HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
FENCE_RE = re.compile(r"^```(\w+)?\s*$")


def ensure_code_comment(code_lines: list[str], lang: str) -> list[str]:
    text = "\n".join(code_lines)
    if "附录注释" in text or "注释" in text:
        return code_lines
    if lang in {"ts", "js", "java"}:
        prefix = "// 附录注释：以下代码摘自项目核心实现，保留关键业务逻辑。"
    elif lang in {"xml", "html"}:
        prefix = "<!-- 附录注释：以下代码摘自项目核心实现，保留关键业务逻辑。 -->"
    else:
        prefix = "# 附录注释：以下代码摘自项目核心实现，保留关键业务逻辑。"
    return [prefix, ""] + code_lines


def add_code_block(doc: Document, code_lines: list[str], lang: str):
    code_lines = ensure_code_comment(code_lines, lang)
    label = doc.add_paragraph(style="SourcePath")
    label.add_run(f"代码片段语言：{lang or 'text'}")

    for line in code_lines:
        p = doc.add_paragraph(style="CodeBlockLine")
        set_paragraph_shading(p, "F7F8FA")
        text = line if line else " "
        run = p.add_run(text)
        set_run_font(run, "Consolas", "SimSun", Pt(8))

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def add_markdown_paragraph(doc: Document, text: str):
    if not text.strip():
        return
    if text.strip() == "---":
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_bottom_border(p)
        return

    style = "SourcePath" if text.startswith("源文件：") else "Normal"
    p = doc.add_paragraph(style=style)
    if text.startswith(">"):
        set_paragraph_shading(p, "F4F6F9")
        text = text.lstrip("> ").strip()
    run = p.add_run(text)
    set_run_font(run, "Calibri", "Microsoft YaHei", Pt(9 if style == "SourcePath" else 11))


def build_docx():
    md = SOURCE_MD.read_text(encoding="utf-8").splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)
    add_footer(doc)
    add_cover(doc)

    in_code = False
    code_lang = ""
    code_lines: list[str] = []

    for raw_line in md:
        fence = FENCE_RE.match(raw_line)
        if fence:
            if not in_code:
                in_code = True
                code_lang = fence.group(1) or "text"
                code_lines = []
            else:
                add_code_block(doc, code_lines, code_lang)
                in_code = False
                code_lang = ""
                code_lines = []
            continue

        if in_code:
            code_lines.append(raw_line)
            continue

        heading = HEADING_RE.match(raw_line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level <= 1:
                p = doc.add_heading(title, level=1)
            elif level == 2:
                p = doc.add_heading(title, level=1)
            elif level == 3:
                p = doc.add_heading(title, level=2)
            else:
                p = doc.add_heading(title, level=3)
            if level == 3 and title[:1].isdigit():
                p.paragraph_format.page_break_before = True
            continue

        add_markdown_paragraph(doc, raw_line)

    if in_code:
        add_code_block(doc, code_lines, code_lang)

    doc.core_properties.title = "BFSM 核心代码实现说明"
    doc.core_properties.subject = "美容美发服务管理系统核心源码摘录与注释"
    doc.core_properties.author = "Codex"
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    path = build_docx()
    print(path)
