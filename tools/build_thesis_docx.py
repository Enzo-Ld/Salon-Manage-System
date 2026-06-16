from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\D\Desktop\毕业设计说明书（论文）模板-计算机、软件.docx")
OUTPUT = ROOT / "docs" / "毕业设计说明书（论文）-李洪东-按模板整理版.docx"


META = {
    "title": "基于Spring Boot的美容美发预约管理系统设计与实现",
    "english_title": "Design and Implementation of a Beauty and Hairdressing Appointment Management System Based on Spring Boot",
    "student": "李洪东",
    "student_no": "220162704018",
    "class_name": "软件工程四班",
    "advisor": "张腾",
    "college": "计算机与人工智能学院",
    "date": "2026年 6 月  日",
}


def clear_body(doc: Document) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_doc_defaults(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(2.6)
    sec.bottom_margin = Cm(2.4)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(21)
    normal.paragraph_format.space_after = Pt(0)

    for name, font, size, bold, color in [
        ("Heading 1", "黑体", 16, True, RGBColor(0, 0, 0)),
        ("Heading 2", "黑体", 14, True, RGBColor(0, 0, 0)),
        ("Heading 3", "宋体", 12, True, RGBColor(0, 0, 0)),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def set_run_east_asia(run, font: str, latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), font)


def configure_footer(doc: Document) -> None:
    sec = doc.sections[0]
    footer = sec.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.add_run("- ")
    add_page_field(p)
    p.add_run(" -")


def add_center(doc: Document, text: str, size: int = 12, bold: bool = False, font: str = "宋体", space_after: int = 6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_run_east_asia(r, font)
    r.font.size = Pt(size)
    r.bold = bold
    return p


def add_para(doc: Document, text: str = "", bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Pt(21)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        set_run_east_asia(r, "黑体")
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_no_indent_para(doc: Document, text: str = "", bold: bool = False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.5
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
    return p


def add_table(doc: Document, caption: str, headers: list[str], rows: Iterable[Iterable[str]], widths: list[float] | None = None):
    cap = add_no_indent_para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_east_asia(cap.runs[0], "楷体")
    cap.runs[0].font.size = Pt(10.5)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, head in enumerate(headers):
        hdr[idx].text = head
        hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in hdr[idx].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                set_run_east_asia(r, "黑体")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[idx].paragraphs:
                p.paragraph_format.first_line_indent = None
                p.paragraph_format.line_spacing = 1.2
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    set_run_east_asia(r, "宋体")
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_code(doc: Document, caption: str, code: str):
    cap = add_no_indent_para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_east_asia(cap.runs[0], "楷体")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for line in code.strip().splitlines():
        r = p.add_run(line.rstrip() + "\n")
        set_run_east_asia(r, "宋体", "Consolas")
        r.font.size = Pt(8)
    return p


def page_break(doc: Document):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_cover(doc: Document):
    add_center(doc, "LANZHOU  UNIVERSITY  OF  TECHNOLOGY", 12, False, "Times New Roman", 24)
    add_center(doc, "毕业设计", 30, True, "黑体", 36)
    add_center(doc, "题目", 14, True, "黑体", 4)
    add_center(doc, META["title"], 18, True, "宋体", 34)
    rows = [
        ("学生姓名", META["student"]),
        ("学    号", META["student_no"]),
        ("专业班级", META["class_name"]),
        ("指导教师", META["advisor"]),
        ("学    院", META["college"]),
        ("答辩日期", META["date"]),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[0].width = Cm(4)
        cells[1].width = Cm(10)
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = None
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(12)
                    set_run_east_asia(r, "宋体")
    page_break(doc)
    add_center(doc, META["title"], 18, True, "宋体", 18)
    add_center(doc, META["english_title"], 14, True, "Times New Roman", 30)
    add_center(doc, f"{META['student']}（LI Hongdong）", 12, False, "宋体", 10)
    add_center(doc, f"学号：{META['student_no']}    专业班级：{META['class_name']}", 12, False, "宋体", 10)
    add_center(doc, f"指导教师：{META['advisor']}", 12, False, "宋体", 10)
    page_break(doc)


def add_abstract(doc: Document):
    add_center(doc, "摘  要", 16, True, "宋体", 18)
    abstract = [
        "随着美容美发行业服务项目不断细化和线上消费习惯逐步形成，传统依靠电话登记、纸质记录和人工统计的管理方式已经难以满足门店预约、排班、支付、订单处理和运营分析等需求。为提高门店业务处理效率和用户服务体验，本文设计并实现了一套基于 Spring Boot 的美容美发预约管理系统。",
        "系统采用前后端分离架构，后端基于 Spring Boot 3.4.4 和 Java 17 构建 REST 服务，结合 Spring Security 与 JWT 实现登录认证和角色权限控制，使用 MyBatis-Plus 操作 MySQL 数据库，并引入 Redis 完成缓存和预约、订单等关键业务的并发控制。前端基于 Vue 3、TypeScript、Vite、Element Plus、Pinia、Vue Router 与 Axios 实现多角色交互界面，使用 ECharts 展示经营统计图表。系统围绕普通用户、技师和管理员三类角色，完成了服务预约、技师推荐、排班管理、请假考勤、商品商城、购物车、商品订单、钱包支付、系统通知、在线聊天、AI 咨询和后台统计等功能。",
        "在核心实现上，预约模块结合服务时长、技师能力、排班、请假和已有预约记录计算可预约性，并通过 Redis 分布式锁避免同一技师同一时间段被重复预约；商城订单模块在下单时校验商品状态和库存，并通过库存锁降低并发超卖风险；支付模块抽象为统一支付单，支持预约支付、商品订单支付和钱包充值，并可扩展到支付宝、微信和手动收银等不同渠道。测试结果表明，系统能够较好地满足美容美发门店线上预约、服务履约、商品销售、支付结算和运营管理的基本需求。",
    ]
    for text in abstract:
        add_para(doc, text)
    add_no_indent_para(doc, "关键词：Spring Boot；美容美发预约；Vue 3；Redis；管理系统", bold=True)
    page_break(doc)

    add_center(doc, "Abstract", 14, True, "Times New Roman", 18)
    english = [
        "With the refinement of beauty and hairdressing services and the popularity of online consumption, traditional management methods based on telephone registration, paper records and manual statistics can no longer meet the needs of appointment scheduling, staff arrangement, payment, order processing and business analysis. To improve operational efficiency and user experience, this thesis designs and implements a beauty and hairdressing appointment management system based on Spring Boot.",
        "The system adopts a front-end and back-end separated architecture. The back end is built with Spring Boot 3.4.4 and Java 17, provides REST services, uses Spring Security and JWT for authentication and role-based authorization, uses MyBatis-Plus to access MySQL, and introduces Redis for caching and concurrency control in key businesses such as appointment creation and order submission. The front end is implemented with Vue 3, TypeScript, Vite, Element Plus, Pinia, Vue Router and Axios, while ECharts is used for statistical visualization. The system supports three roles: users, technicians and administrators, and implements service appointment, technician recommendation, scheduling, leave and attendance, mall orders, wallet payment, notifications, online chat, AI consultation and management statistics.",
        "For the core business, the appointment module checks service duration, technician capability, schedules, leave records and existing appointments, and uses Redis distributed locks to prevent duplicate appointments. The mall order module validates product status and stock, and uses stock locking to reduce overselling risks. The payment module is modeled as a unified payment order and supports appointment payment, product order payment and wallet recharge, with extensible payment gateways. Test results show that the system can meet the basic needs of online appointment, service fulfillment, product sales, payment settlement and operation management in beauty and hairdressing stores.",
    ]
    for text in english:
        add_para(doc, text)
    add_no_indent_para(doc, "Key words: spring boot; beauty and hairdressing appointment; vue 3; redis; management system", bold=True)
    page_break(doc)


def add_toc(doc: Document):
    add_center(doc, "目  录", 16, True, "宋体", 18)
    entries = [
        "摘  要",
        "Abstract",
        "第1章  绪论",
        "1.1 研究背景及意义",
        "1.2 国内外的发展状况及应用",
        "1.3 毕业设计主要内容及说明书的组织结构",
        "第2章  系统需求分析",
        "2.1 功能需求",
        "2.2 性能需求",
        "2.3 技术可行性分析",
        "2.4 非技术经济因素和工程伦理",
        "第3章  系统总体设计",
        "3.1 系统体系架构设计",
        "3.2 系统功能结构设计",
        "3.3 系统核心算法设计",
        "3.4 系统数据库设计",
        "3.5 实现系统的关键技术",
        "第4章  系统详细设计及实现",
        "4.1 认证与权限模块的设计及实现",
        "4.2 预约与排班模块的设计及实现",
        "4.3 商城订单与支付模块的设计及实现",
        "4.4 消息通知与 AI 咨询模块的设计及实现",
        "4.5 统计与后台管理模块的设计及实现",
        "第5章  系统测试及使用说明",
        "5.1 系统测试报告",
        "5.2 系统使用说明",
        "设计总结",
        "参考文献",
        "外文翻译",
        "附录  程序源代码",
        "致  谢",
    ]
    for e in entries:
        p = add_no_indent_para(doc, e)
        if e.startswith("第") or e in {"设计总结", "参考文献", "外文翻译", "附录  程序源代码", "致  谢"}:
            p.runs[0].bold = True
    page_break(doc)


def add_chapter1(doc: Document):
    add_heading(doc, "第1章  绪论", 1)
    add_heading(doc, "1.1 研究背景及意义", 2)
    add_heading(doc, "1.1.1 研究背景", 3)
    add_para(doc, "近年来，美容美发行业的服务内容由传统理发逐步扩展到发型设计、头皮养护、面部护理、身体理疗、护理产品销售等多种业务。用户对服务质量、预约效率和消费体验的要求不断提高，门店也需要对技师排班、服务履约、商品库存、订单支付和客户评价进行更精细化管理。传统电话预约、纸质登记和简单表格记录方式容易出现预约遗漏、排班冲突、订单状态不清和统计滞后等问题，难以支撑多角色协同和线上线下一体化经营。")
    add_para(doc, "Web 技术的发展为中小型服务门店提供了低成本、易维护的信息化方案。Spring Boot 能够快速构建稳定的后端服务，Vue 3 能够提供良好的前端交互体验，MySQL 与 Redis 能够分别承担业务数据持久化和关键业务并发控制。本课题以美容美发门店实际经营流程为背景，设计并实现一套集服务预约、技师管理、商品商城、支付结算、消息沟通、AI 咨询和统计分析于一体的管理系统。")
    add_heading(doc, "1.1.2 研究意义", 3)
    add_para(doc, "本系统的研究意义主要体现在三个方面。第一，在管理效率方面，系统将用户、技师、服务项目、排班、预约、商品、订单和支付等数据统一管理，减少人工登记和人工核对造成的错误。第二，在用户体验方面，用户能够在线浏览服务和商品，完成预约、支付、评价、收藏和咨询，降低线下沟通成本。第三，在经营决策方面，系统通过销售趋势、商品排行、分类统计、技师业绩和考勤图表，为管理员提供门店运营参考。")
    add_heading(doc, "1.2 国内外的发展状况及应用", 2)
    add_heading(doc, "1.2.1 国内外发展状况", 3)
    add_para(doc, "国内外本地生活服务平台已经普遍支持线上预约、在线支付、评价反馈和商家管理等能力，大型平台更强调流量聚合和标准化服务，中小型门店则更关注低成本部署、数据自主维护和业务流程适配。美容美发行业的管理系统通常从会员管理、门店预约、员工排班、消费记录和客户评价等功能展开，部分系统开始引入智能推荐和数据分析能力。")
    add_para(doc, "从技术趋势看，前后端分离、REST API、JWT 无状态认证、组件化前端、缓存与分布式锁、数据可视化等技术已经成为管理系统开发中的常见方案。与传统单体页面应用相比，前后端分离方式更适合多端扩展和角色差异化界面构建，也便于后端接口复用和业务模块迭代。")
    add_heading(doc, "1.2.2 美容美发预约管理系统的应用", 3)
    add_para(doc, "美容美发预约管理系统主要应用于具有到店服务、技师资源和预约时段约束的门店场景。系统通过线上服务展示、可预约时间计算、技师推荐、预约支付和服务评价形成服务闭环；通过商品商城、购物车、订单发货和售后退款形成商品销售闭环；通过请假、考勤、消息通知和后台统计形成门店内部管理闭环。本课题在实现中还引入 AI 形象顾问模块，将用户咨询与系统内服务、商品目录结合，提升系统的智能化和互动性。")
    add_heading(doc, "1.3 毕业设计主要内容及说明书的组织结构", 2)
    add_heading(doc, "1.3.1 毕业设计主要内容", 3)
    add_para(doc, "本文围绕普通用户、技师和管理员三类角色完成系统设计与实现。用户端实现注册登录、服务浏览、预约创建、支付、商品购买、钱包充值、评价收藏、通知查看、聊天沟通和 AI 咨询；技师端实现工作台、预约处理、请假申请、考勤打卡和个人统计；管理员端实现用户管理、技师与服务维护、排班管理、预约管理、商品与订单管理、评论管理、轮播资讯、充值活动和经营统计。")
    add_heading(doc, "1.3.2 说明书的组织结构", 3)
    add_para(doc, "第1章介绍课题背景、意义和研究内容；第2章分析系统功能需求、性能需求、可行性和工程伦理；第3章给出系统总体架构、功能结构、核心算法、数据库设计和关键技术；第4章结合项目代码说明主要模块的详细设计与实现；第5章说明系统测试过程、测试结果和使用方法；最后给出设计总结、参考文献、外文翻译、附录和致谢。")
    page_break(doc)


def add_chapter2(doc: Document):
    add_heading(doc, "第2章  系统需求分析", 1)
    add_heading(doc, "2.1 功能需求", 2)
    add_para(doc, "系统需求围绕普通用户、技师和管理员三类角色展开。普通用户需要完成服务浏览、在线预约、预约支付、商品浏览、购物车、商品订单、钱包充值、个人资料与头像维护、收货地址管理、收藏记录、评论评价、站内通知、在线聊天和 AI 咨询等操作。技师需要查看与处理预约，按照业务规则推进服务状态，提交请假申请，完成上下班打卡，并查看个人业绩和考勤统计。管理员需要维护用户、角色、技师、服务、排班、预约、评论、商品、商品分类、订单、售后、轮播图、资讯和充值活动，并查看经营统计。")
    add_table(doc, "表2-1 系统角色与主要功能需求", ["角色", "主要功能"], [
        ["普通用户", "注册登录、服务预约、预约支付、商品购买、钱包充值、评价收藏、通知查看、聊天和 AI 咨询"],
        ["技师", "预约查看、状态处理、请假申请、考勤打卡、个人业绩和考勤统计"],
        ["管理员", "用户管理、服务和技师管理、排班管理、预约管理、商品订单管理、评论管理、内容管理和统计分析"],
    ], [3, 12])
    add_para(doc, "从业务流程看，服务预约是系统的核心流程。用户选择服务项目和预约时间后，系统根据服务时长计算结束时间，结合技师能力、排班、请假和已有预约记录判断是否可预约；如果用户未指定技师，系统自动推荐合适技师；预约创建后需要支付，支付成功后进入待确认和服务履约流程。")
    add_para(doc, "商城订单流程以商品浏览、购物车、收货地址、订单创建、支付、发货、确认收货和售后退款为主。支付流程抽象为统一支付单，既服务商品订单，也服务预约订单和钱包充值，使支付状态、业务单状态和钱包流水之间保持一致。")
    add_heading(doc, "2.2 性能需求", 2)
    add_para(doc, "系统应满足中小型美容美发门店日常业务访问需求。在并发预约和商品下单场景中，系统需要避免同一技师同一时段重复预约以及商品库存超卖。对于服务列表、商品列表、订单列表、通知列表和统计数据，系统需要支持分页查询和条件筛选，保证页面响应及时。前端页面应具备清晰的角色入口和状态提示，减少误操作。")
    add_para(doc, "数据一致性也是重要性能目标。预约、订单、支付和退款等业务涉及多个数据表的状态变更，应在事务范围内完成，并在关键资源上使用 Redis 分布式锁控制并发。对于异常输入、未登录访问、越权访问、非法状态流转等情况，系统应返回统一错误结构并给出明确提示。")
    add_heading(doc, "2.3 技术可行性分析", 2)
    add_para(doc, "本系统所采用的技术均为成熟开源技术。Spring Boot 3.4.4 能够快速集成 Web、Security、Validation、Redis、WebSocket 等能力；MyBatis-Plus 可以简化数据库增删改查和分页操作；MySQL 适合存储结构化业务数据；Redis 可用于缓存和分布式锁；Vue 3、Element Plus 和 ECharts 能够支撑多角色前端页面和可视化统计。项目源码中已经实现 Controller、Service、Mapper、Entity、DTO、VO 等分层结构，具备完整的开发基础。")
    add_table(doc, "表2-2 系统主要技术与用途", ["技术", "用途"], [
        ["Spring Boot 3.4.4", "构建后端 REST 服务和业务处理框架"],
        ["Spring Security + JWT", "实现登录认证、无状态会话和角色权限控制"],
        ["MyBatis-Plus + MySQL", "实现业务数据持久化、分页和条件查询"],
        ["Redis", "缓存热点数据，提供预约和库存并发锁"],
        ["Vue 3 + Element Plus", "构建用户端、技师端和管理端页面"],
        ["ECharts + WebSocket", "展示统计图表，实现聊天消息推送"],
    ], [4, 11])
    add_heading(doc, "2.4 非技术经济因素和工程伦理", 2)
    add_para(doc, "从经济因素看，本系统使用 Java、Spring Boot、Vue、MySQL、Redis 等开源技术，开发与部署成本较低，适合毕业设计和中小型门店原型系统建设。系统模块划分清晰，后续可按门店需求扩展会员、优惠券、短信提醒、真实支付回调等功能。")
    add_para(doc, "从工程伦理看，系统需要保护用户账号、手机号、头像、收货地址、预约记录和聊天记录等个人信息。实现中通过密码加密、JWT 鉴权、角色权限控制和统一异常处理降低数据泄露风险。AI 咨询功能仅作为辅助建议，不应替代专业美容美发服务判断；系统应避免夸大 AI 推荐效果，并在后续实际应用中增加隐私授权和敏感内容过滤。")
    page_break(doc)


def add_chapter3(doc: Document):
    add_heading(doc, "第3章  系统总体设计", 1)
    add_heading(doc, "3.1 系统体系架构设计", 2)
    add_para(doc, "系统采用 B/S 架构和前后端分离设计。前端位于 `frontend` 目录，负责页面展示、路由控制、状态管理、请求封装和图表展示；后端位于 `backend` 目录，负责认证授权、业务规则、数据持久化、缓存并发控制、文件上传、支付抽象和 WebSocket 消息推送；数据库脚本位于 `backend/src/main/resources/schema.sql` 和 `bfsm.sql`。")
    add_table(doc, "表3-1 系统分层结构", ["层次", "主要组件", "职责"], [
        ["表现层", "Vue 3、Element Plus、ECharts", "多角色页面展示、表单交互、图表展示和路由控制"],
        ["接口层", "Controller、ApiResponse", "接收 REST 请求、参数绑定和统一响应"],
        ["业务层", "Service、Gateway、RedisLockService", "处理预约、支付、订单、排班、通知、AI 咨询等业务规则"],
        ["数据层", "Mapper、MyBatis-Plus、MySQL", "完成实体映射、分页查询和数据持久化"],
        ["支撑层", "Spring Security、JWT、Redis、WebSocket", "认证鉴权、缓存锁、实时消息和静态资源访问"],
    ], [2.5, 5, 7.5])
    add_no_indent_para(doc, "图3-1 系统总体调用流程", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "用户在 Vue 页面发起操作后，Axios 请求后端 REST API；Spring Security 通过 JWT 识别用户身份；Controller 将请求转交给 Service；Service 调用 Mapper 访问 MySQL，并在必要时使用 Redis 锁、WebSocket 和通知服务；最后通过统一 ApiResponse 返回前端。")
    add_heading(doc, "3.2 系统功能结构设计", 2)
    add_para(doc, "系统功能结构按照角色和业务线划分。用户端重点完成消费闭环，技师端重点完成服务履约，管理员端重点完成门店运营配置和数据分析。")
    add_table(doc, "表3-2 系统功能模块划分", ["模块", "包含功能"], [
        ["认证与权限", "注册、登录、JWT 鉴权、角色路由、个人资料、头像上传"],
        ["预约与排班", "服务浏览、可约时段、技师推荐、预约创建、预约支付、状态流转、评价"],
        ["请假考勤", "技师请假、管理员审批、冲突处理、上下班打卡、月度统计"],
        ["商城订单", "商品分类、商品详情、购物车、订单创建、发货、收货、售后退款"],
        ["钱包支付", "统一支付单、余额支付、钱包充值、充值活动、退款记录"],
        ["互动通知", "系统通知、收藏、评论、在线聊天、WebSocket 推送"],
        ["AI 咨询", "图片上传、咨询会话、服务和商品推荐、历史会话"],
        ["后台统计", "预约、销售、商品排行、分类统计、技师业绩和考勤图表"],
    ], [3, 12])
    add_heading(doc, "3.3 系统核心算法设计", 2)
    add_para(doc, "系统核心算法主要体现在技师推荐、可预约时段计算和并发控制。技师推荐由 `SchedulingService` 实现，先根据服务项目查询具备对应能力的在职技师，再过滤排班不覆盖、请假冲突和已有预约冲突的技师，最后根据评分和当天负载计算推荐分数。")
    add_no_indent_para(doc, "推荐分数计算公式如下：", bold=True)
    add_center(doc, "score = 60 + 技师评分 × 8 - 当天负载 × 10", 11, False, "Times New Roman", 6)
    add_para(doc, "该公式使高评分技师具有更高推荐优先级，同时通过当天负载扣分实现简单的均衡分配。可预约时段计算以服务项目时长为基础，将预约开始时间和结束时间映射到技师排班范围内，并排除已请假和已存在有效预约的区间。预约创建和商品下单等关键操作使用 Redis 分布式锁保护，避免并发请求同时占用同一资源。")
    add_heading(doc, "3.4 系统数据库设计", 2)
    add_para(doc, "数据库围绕用户权限、门店基础资料、排班预约、商城订单、支付钱包、互动通知、聊天和 AI 咨询等业务对象设计。核心表均在 `schema.sql` 和 `bfsm.sql` 中定义，实体类位于 `backend/src/main/java/com/bfsm/entity`。")
    add_table(doc, "表3-3 数据库核心表分组", ["业务分组", "数据表"], [
        ["用户权限", "sys_user、sys_role、sys_user_role"],
        ["基础资料", "technician、technician_service、service_category、service_item、banner、announcement"],
        ["排班预约", "schedule、leave_record、attendance_record、appointment、appointment_log"],
        ["商城订单", "product_category、product、shopping_cart_item、user_address、product_order、product_order_item"],
        ["支付钱包", "user_wallet、wallet_txn、recharge_activity、payment_order、payment_refund"],
        ["互动通知", "favorite、like_record、comment、system_notification、chat_session、chat_message"],
        ["AI 咨询", "ai_consult_record、ai_consult_message"],
    ], [3.5, 11.5])
    add_para(doc, "其中，`appointment` 保存预约主数据，包含用户、技师、服务项目、预约时间、结束时间、金额、业务状态和支付状态；`schedule` 保存技师排班；`leave_record` 保存请假记录；`product_order` 和 `product_order_item` 保存商品订单及明细；`payment_order` 统一记录预约、商品订单和钱包充值的支付行为；`wallet_txn` 记录钱包余额变化。")
    add_heading(doc, "3.5 实现系统的关键技术", 2)
    add_para(doc, "系统关键技术包括 JWT 无状态认证、统一响应与异常处理、MyBatis-Plus 数据访问、Redis 分布式锁、支付网关抽象、WebSocket 消息推送和 AI 兼容接口。后端通过 `JwtAuthFilter` 解析请求头中的 token，将用户身份写入 Spring Security 上下文；通过 `ApiResponse` 和 `GlobalExceptionHandler` 统一成功与失败响应；通过 `PaymentGatewayRouter` 将不同支付方式路由到支付宝、微信或手动收银实现；通过 `ChatWebSocketHandler` 维护在线连接并推送聊天消息。")


def add_chapter4(doc: Document):
    add_heading(doc, "第4章  系统详细设计及实现", 1)
    add_heading(doc, "4.1 认证与权限模块的设计及实现", 2)
    add_heading(doc, "4.1.1 认证与权限模块的用户界面设计", 3)
    add_para(doc, "认证与权限模块面向所有角色提供登录、注册、个人资料和头像维护功能。前端使用 Vue Router 根据当前用户角色控制页面入口，普通用户进入服务、商城、钱包和 AI 咨询相关页面，技师进入工作台、预约处理和请假考勤页面，管理员进入后台管理和统计页面。Axios 请求拦截器会自动携带 JWT，未登录或 token 失效时跳转登录页。")
    add_heading(doc, "4.1.2 认证与权限模块的程序处理流程设计", 3)
    add_para(doc, "用户注册时，后端只允许公开注册普通用户角色，密码经 BCrypt 加密后保存。登录时，`AuthService` 校验用户名、账号启用状态和密码，校验通过后由 `JwtTokenService` 生成 token。后续请求经 `JwtAuthFilter` 解析 token，并将用户 ID、用户名和角色放入安全上下文，业务层通过 `SecurityUtils.currentUserOrThrow()` 获取当前用户。")
    add_code(doc, "程序4-1 JWT 认证过滤流程片段", """
String token = resolveToken(request);
if (token != null && jwtTokenService.validate(token)) {
    SecurityUser user = jwtTokenService.parse(token);
    Authentication auth = new UsernamePasswordAuthenticationToken(
            user, null, user.getAuthorities());
    SecurityContextHolder.getContext().setAuthentication(auth);
}
filterChain.doFilter(request, response);
""")
    add_heading(doc, "4.2 预约与排班模块的设计及实现", 2)
    add_heading(doc, "4.2.1 预约与排班模块的用户界面设计", 3)
    add_para(doc, "用户端服务页面展示服务分类、服务项目、价格、时长、评价和收藏状态。用户选择服务后进入预约窗口，选择日期、时段和技师；如果不指定技师，系统自动推荐。技师端预约页面用于查看待确认、待服务、服务中和已完成预约，并按业务规则推进状态。管理员端预约与排班页面用于维护技师班次、查看预约列表、处理改派和审核。")
    add_heading(doc, "4.2.2 预约与排班模块的程序处理流程设计", 3)
    add_para(doc, "预约创建由 `AppointmentService.create` 实现。方法首先校验预约时间不能早于当前时间，再读取服务项目并根据服务时长计算结束时间。如果请求未指定技师，则调用 `SchedulingService.recommendTechnician` 推荐技师；如果指定技师，则调用 `ScheduleConflictChecker` 校验服务能力、排班、请假和已有预约冲突。随后系统使用 Redis 锁锁定技师与预约时间，在锁内再次校验冲突并创建预约单。")
    add_code(doc, "程序4-2 预约创建核心逻辑片段", """
if (req.getAppointmentTime().isBefore(LocalDateTime.now())) {
    throw new BizException(4010, "预约时间不能早于当前时间");
}
ServiceItem item = serviceItemOrThrow(req.getServiceItemId());
LocalDateTime endTime =
        req.getAppointmentTime().plusMinutes(item.getDurationMin());
Long technicianId = req.getTechnicianId();
if (technicianId == null) {
    technicianId = schedulingService.recommendTechnician(
            req.getServiceItemId(), req.getAppointmentTime())
            .getTechnicianId();
}
String token = redisLockService.tryLock(lockKey, Duration.ofSeconds(8));
try {
    scheduleConflictChecker.validateTechnicianForAppointment(
            technicianId, req.getServiceItemId(), req.getAppointmentTime(), endTime);
    Appointment appointment = new Appointment();
    appointment.setStatus("CREATED");
    appointment.setPayStatus("UNPAID");
    appointmentMapper.insert(appointment);
} finally {
    redisLockService.unlock(lockKey, token);
}
""")
    add_para(doc, "技师推荐由 `SchedulingService` 实现，核心是先过滤不可服务、未排班、请假和时间冲突的技师，再按照推荐分数选择最优候选。预约状态流转由 `FLOW` 常量约束，只允许从待确认到已确认或拒绝、从已确认到待服务、从待服务到服务中、从服务中到已完成，避免非法跳转。")
    add_code(doc, "程序4-3 技师推荐筛选与评分片段", """
return candidates.stream()
    .filter(t -> scheduleConflictChecker.isWorkingDuring(t.getId(), start, end))
    .filter(t -> !scheduleConflictChecker.isOnLeaveDuring(t.getId(), start, end))
    .filter(t -> !scheduleConflictChecker.hasConflictInRange(t.getId(), start, end))
    .map(t -> toRecommendation(t, start, end, item.getDurationMin()))
    .max(Comparator.comparingDouble(TechnicianRecommendationVo::getScore))
    .orElse(null);
double score = 60 + rating * 8 - dayLoad * 10;
""")
    add_heading(doc, "4.3 商城订单与支付模块的设计及实现", 2)
    add_heading(doc, "4.3.1 商城订单与支付模块的用户界面设计", 3)
    add_para(doc, "商城页面展示商品分类、商品列表、商品详情、价格、库存、收藏和评论。用户可以将商品加入购物车，也可以直接下单。订单页面展示待支付、待发货、待收货、已完成和售后状态。支付收银台根据后端返回的支付能力展示钱包余额、微信、支付宝或手动收银入口，毕业设计演示环境可通过手动收银完成支付状态流转。")
    add_heading(doc, "4.3.2 商城订单与支付模块的程序处理流程设计", 3)
    add_para(doc, "商品订单由 `ProductOrderService` 处理。创建订单时，系统校验商品是否上架、库存是否充足，使用 Redis 锁保护库存扣减，生成 `product_order` 与 `product_order_item`，并在来自购物车时删除对应购物车项。支付由 `PaymentService` 统一处理，预约订单、商品订单和钱包充值均会生成 `payment_order`，支付成功后同步更新业务单状态、钱包流水和通知。")
    add_code(doc, "程序4-4 统一支付单创建逻辑片段", """
closePendingPayments("PRODUCT_ORDER", order.getId(), "A new payment was created");
PaymentOrder paymentOrder = new PaymentOrder();
paymentOrder.setBusinessType("PRODUCT_ORDER");
paymentOrder.setBusinessId(order.getId());
paymentOrder.setPaymentNo(buildPaymentNo("P"));
paymentOrder.setPayChannel(payChannel);
paymentOrder.setPayMode(resolveSupportedPayMode(payChannel));
paymentOrder.setAmount(order.getPayAmount());
paymentOrder.setStatus("PAYING");
paymentOrder.setExpireTime(LocalDateTime.now().plusMinutes(15));
paymentOrderMapper.insert(paymentOrder);
return gatewayForMode(paymentOrder.getPayMode())
        .create(paymentOrder);
""")
    add_para(doc, "支付模块没有将支付逻辑写死在预约或商品订单中，而是通过统一支付单和支付网关路由实现复用。`ManualPaymentGateway` 用于演示环境，`AlipayPagePaymentGateway` 和 `WechatNativePaymentGateway` 可对接真实支付配置。定时任务会关闭过期支付单，退款业务通过 `payment_refund` 记录。")
    add_heading(doc, "4.4 消息通知与 AI 咨询模块的设计及实现", 2)
    add_heading(doc, "4.4.1 消息通知与 AI 咨询模块的用户界面设计", 3)
    add_para(doc, "通知中心展示业务通知、未读数量、已读状态和清理入口。聊天页面展示会话列表和消息记录，用户与技师可以围绕预约服务进行沟通。AI 咨询页面支持上传图片和输入发型、护理、预算或购买意图，系统返回发型建议、护理建议、服务推荐和商品推荐。")
    add_heading(doc, "4.4.2 消息通知与 AI 咨询模块的程序处理流程设计", 3)
    add_para(doc, "通知由 `NotificationService` 在预约、订单、请假、售后等关键业务节点生成。聊天由 `ChatService` 保存会话和消息，再通过 `ChatWebSocketHandler` 推送给在线接收者。AI 咨询由 `AiConsultService` 调用兼容式大模型客户端，结合系统服务目录和商品目录构造上下文，保存咨询记录和推荐快照；当外部接口不可用时，可使用 Mock 客户端或本地降级策略保证功能可演示。")
    add_code(doc, "程序4-5 AI 咨询上下文构造示意", """
AiConsultContext context = new AiConsultContext();
context.setPrompt(req.getPrompt());
context.setImageUrl(req.getImageUrl());
context.setServices(loadEnabledServiceCatalog());
context.setProducts(loadOnSaleProductCatalog());
AiBeautyAnalysis result = advisorClient.analyze(context);
saveConsultRecord(currentUserId, req, result);
return toConsultVo(result);
""")
    add_heading(doc, "4.5 统计与后台管理模块的设计及实现", 2)
    add_heading(doc, "4.5.1 统计与后台管理模块的用户界面设计", 3)
    add_para(doc, "后台管理页面按照业务域划分为用户管理、服务运营、排班管理、预约管理、商品运营、订单管理、评论管理、内容管理、充值活动和统计看板。Element Plus 表格、表单、弹窗和分页组件用于提高管理效率。统计页面通过 ECharts 展示销售额趋势、商品销量排行、分类统计、技师服务量、评分、业绩和考勤数据。")
    add_heading(doc, "4.5.2 统计与后台管理模块的程序处理流程设计", 3)
    add_para(doc, "后台接口主要由 `AdminController`、`MasterDataController`、`ProductController` 和 `StatsController` 提供。统计服务 `StatsService` 通过订单、预约、商品和考勤数据计算概览指标、趋势数据和排行数据，并以 VO 对象返回前端。用户管理支持分页查询、账号启停、角色调整和密码重置；排班管理支持手动排班、自动排班预览和应用，并在取消排班时检查是否已有有效预约。")
    add_table(doc, "表4-1 主要后端接口分组", ["接口前缀", "功能说明"], [
        ["/api/auth", "注册、登录、当前用户、个人资料、头像上传"],
        ["/api/appointments", "预约创建、分页、详情、技师推荐、可约时段、状态流转、取消和评价"],
        ["/api/products", "商品分类、商品列表、购物车、地址、订单、发货和售后"],
        ["/api/payments", "支付单创建、支付状态、手动完成、退款处理和支付回调"],
        ["/api/master", "技师、服务、轮播、资讯和排班维护"],
        ["/api/stats", "概览、趋势、服务统计、技师统计、商城销售统计"],
    ], [4, 11])


def add_chapter5(doc: Document):
    add_heading(doc, "第5章  系统测试及使用说明", 1)
    add_heading(doc, "5.1 系统测试报告", 2)
    add_heading(doc, "5.1.1 测试的目的及原则", 3)
    add_para(doc, "系统测试的目的在于验证功能完整性、业务流程正确性、角色权限安全性和异常场景处理能力。本系统采用黑盒测试与白盒测试相结合的方法：黑盒测试从页面交互、接口返回和业务状态变化角度验证功能；白盒测试结合后端 JUnit5 与 Mockito 单元测试验证服务层关键逻辑。")
    add_para(doc, "测试原则包括：核心业务覆盖正常流程和异常流程；涉及状态流转的功能验证合法和非法跳转；涉及数据变更的功能同时检查页面反馈、接口响应和数据库状态；涉及权限的功能验证未登录、角色不匹配和账号禁用等场景。")
    add_heading(doc, "5.1.2 测试过程及结果", 3)
    add_table(doc, "表5-1 系统测试环境", ["测试项", "测试环境"], [
        ["操作系统", "Windows PC"],
        ["浏览器", "Chrome 浏览器"],
        ["接口测试工具", "Postman"],
        ["后端运行环境", "JDK 17、Maven"],
        ["前端运行环境", "Node.js、Vite"],
        ["后端框架", "Spring Boot 3.4.4"],
        ["前端框架", "Vue3、TypeScript、Element Plus"],
        ["数据库", "MySQL 8"],
        ["缓存与并发控制", "Redis"],
        ["端口", "后端 8080，前端 5173"],
    ], [4, 11])
    add_table(doc, "表5-2 系统测试账号", ["账号", "密码", "角色", "用途"], [
        ["admin", "123456", "管理员", "后台管理、统计、审批、发货"],
        ["tech_amy", "123456", "技师", "预约处理、请假、考勤"],
        ["tech_bob", "123456", "技师", "技师推荐、排班测试"],
        ["user_tom", "123456", "普通用户", "预约、商城、支付、AI 咨询"],
        ["user_lily", "123456", "普通用户", "多用户业务验证"],
    ], [3, 3, 3, 6])
    add_table(doc, "表5-3 主要测试用例与结果", ["模块", "测试内容", "预期结果", "实际结果"], [
        ["认证与权限", "正确登录、错误密码、越权访问、非法角色注册、头像格式校验", "返回 token 或明确错误提示，路由和接口均能拦截越权访问", "与预期一致"],
        ["服务预约", "服务查询、可约时段、技师推荐、创建预约、取消、评价和非法状态流转", "预约状态和支付状态正确变化，异常场景被拒绝", "与预期一致"],
        ["排班请假", "新增排班、重叠排班、取消已有预约排班、请假审批、考勤打卡", "系统避免排班与预约、请假之间产生冲突", "与预期一致"],
        ["商城订单", "商品查询、购物车、创建订单、库存不足、支付、发货、确认收货和售后", "订单状态按 CREATED、PAID、SHIPPED、COMPLETED 等流程变化", "与预期一致"],
        ["支付钱包", "钱包充值、余额支付、重复支付、余额不足、退款记录", "支付单、业务单和钱包流水保持一致", "与预期一致"],
        ["通知聊天 AI", "通知未读、聊天消息、AI 图片上传、服务推荐、商品推荐和历史会话删除", "业务提醒可查看，聊天可落库，AI 咨询可返回推荐", "与预期一致"],
        ["后台统计", "管理员看板、销售趋势、商品排行、分类统计和技师月度统计", "图表数据能够正常展示", "与预期一致"],
    ], [3, 6, 4, 3])
    add_para(doc, "后端服务层已包含 `AppointmentServiceTest`、`SchedulingServiceTest`、`ProductOrderServiceTest`、`PaymentServiceTest`、`LeaveRecordServiceTest`、`AttendanceServiceTest`、`NotificationServiceTest`、`AiConsultServiceTest` 等单元测试，覆盖预约冲突、状态流转、支付入账、库存扣减、请假审批、考勤限制和 AI 推荐意图等逻辑。")
    add_heading(doc, "5.1.3 测试总结", 3)
    add_para(doc, "测试结果表明，系统能够完成普通用户、技师和管理员三类角色的主要业务流程。对于错误密码、越权访问、非法注册角色、非法头像格式、过去时间预约、未排班预约、预约冲突、重叠排班、库存不足、重复支付、余额不足和请假冲突等异常场景，系统均能给出明确提示并阻止非法操作。")
    add_heading(doc, "5.2 系统使用说明", 2)
    add_heading(doc, "5.2.1 系统简介", 3)
    add_para(doc, "本系统是面向美容美发门店的综合预约管理平台。普通用户通过系统完成服务预约、商品购买和咨询互动；技师通过系统处理预约、请假和考勤；管理员通过后台维护基础数据、处理订单和查看经营统计。")
    add_heading(doc, "5.2.2 系统运行环境", 3)
    add_para(doc, "后端运行需要 JDK 17、Maven、MySQL 8 和 Redis，主要配置位于 `backend/src/main/resources/application.yml`。前端运行需要 Node.js 和 npm，主要依赖位于 `frontend/package.json`。后端默认端口为 8080，前端开发端口为 5173。")
    add_heading(doc, "5.2.3 系统操作说明", 3)
    add_para(doc, "管理员登录后进入后台首页，可维护用户、服务、技师、排班、商品、订单、评论、资讯和轮播图，并查看销售统计。普通用户登录后进入用户端首页，可浏览服务和商品，创建预约或订单并完成支付。技师登录后进入工作台，可处理预约状态、提交请假、完成上下班打卡并查看个人统计。")


def add_conclusion_refs_translation_appendix(doc: Document):
    add_heading(doc, "设计总结", 1)
    add_para(doc, "本文完成了基于 Spring Boot 的美容美发预约管理系统的设计与实现。系统采用前后端分离架构，将普通用户、技师和管理员三类角色的业务需求统一到一个平台中，围绕服务预约、商品商城、钱包支付、排班请假、通知聊天、AI 咨询和后台统计构建较完整的业务闭环。")
    add_para(doc, "在实现过程中，预约模块结合排班、请假、服务时长和预约冲突进行可预约性判断，并通过 Redis 分布式锁解决并发预约问题；商城模块通过库存校验和库存锁降低超卖风险；支付模块抽象为统一支付单，增强了预约、订单和钱包充值场景的复用性；后台统计模块将业务数据转化为可视化图表，为门店经营提供参考。")
    add_para(doc, "本系统仍有继续完善空间，例如可增加真实短信提醒、更加精细的会员积分、优惠券营销、真实支付证书配置、AI 咨询隐私授权和移动端适配等功能。总体来看，系统达到了毕业设计对需求分析、总体设计、详细实现和测试验证的要求。")
    page_break(doc)

    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] 黑马程序员. Spring Boot企业级开发教程[M]. 北京: 人民邮电出版社, 2024.",
        "[2] 趣千厘. Vue.js 3.x+Element Plus前端开发实战[M]. 北京: 清华大学出版社, 2022.",
        "[3] 王彦懿. 媒介使用促进美容美发从业者职业适应的质性研究[D]. 重庆大学, 2022.",
        "[4] 林杰. 基于.net的美容美发管理系统[D]. 吉林大学, 2016.",
        "[5] 陈谦民, 高越, 叶益成, 等. 基于Java的美发管理系统的设计与应用[J]. 现代信息科技, 2022, 6(07):1-7.",
        "[6] 翟晓猛, 陈祯曌, 彭志杰, 等. 一款美发预约APP的设计和研究[J]. 湖北农机化, 2021(02):110-111.",
        "[7] 刘芳. 美发店客户关系管理系统设计与实现[J]. 经营与管理, 2017(08):123-125.",
        "[8] 王珊, 萨师煊. 数据库系统概论: 第6版[M]. 北京: 高等教育出版社, 2022.",
        "[9] 李刚. 轻量级Java EE企业应用实战[M]. 北京: 电子工业出版社, 2023.",
        "[10] 陈雄华. Spring Boot 2.7企业级应用开发实战[M]. 北京: 机械工业出版社, 2023.",
        "[11] 王颖洁. 基于微信小程序的理发店管理系统的设计与实现[J]. 软件, 2024, 45(6):22-27.",
        "[12] 张艳. 基于Spring Boot的美容服务预约系统设计[J]. 电脑知识与技术, 2024, 20(12):67-69.",
        "[13] 刘畅. 基于Spring Boot+Vue的理发店管理系统设计与实现[D]. 济南: 山东师范大学, 2023.",
        "[14] 苏伟鸿. 基于劝导理论的本地生活美发服务平台设计研究与实践[D]. 长沙: 湖南大学, 2023.",
        "[15] Sommerville I. Software Engineering[M]. 10th ed. Boston: Pearson, 2016.",
        "[16] Gamma E, Helm R, Johnson R, et al. Design Patterns: Elements of Reusable Object-Oriented Software[M]. Boston: Addison-Wesley, 1995.",
        "[17] Walls C. Spring Boot in Action[M]. New York: Manning Publications, 2016.",
        "[18] Freeman E, Robson E. Head First Design Patterns[M]. Sebastopol: O'Reilly Media, 2021.",
        "[19] Fowler M. Patterns of Enterprise Application Architecture[M]. Boston: Addison-Wesley, 2003.",
        "[20] Fielding R T. Architectural Styles and the Design of Network-based Software Architectures[D]. University of California, Irvine, 2000.",
    ]
    for ref in refs:
        add_no_indent_para(doc, ref)
    page_break(doc)

    add_heading(doc, "外文翻译", 1)
    add_heading(doc, "原文", 2)
    add_para(doc, "A web-based management system should separate presentation concerns from business rules and data persistence. In a layered architecture, the client is responsible for user interaction, the controller receives requests, the service layer enforces business rules, and the persistence layer stores structured data. This separation improves maintainability because changes in one layer usually have limited influence on other layers.")
    add_para(doc, "For appointment-oriented systems, consistency is more important than simply inserting records. The application must check resource availability, time conflicts, user permissions and payment status before confirming an appointment. When multiple users operate concurrently, locks or database constraints are required to protect shared resources such as staff schedules and product stock.")
    add_para(doc, "Security is also a basic requirement of modern web applications. Passwords should be stored in encrypted form, requests should be authenticated, and users should only access resources that match their roles. For front-end and back-end separated systems, token-based authentication is commonly used to keep the server stateless and to simplify cross-platform access.")
    add_heading(doc, "译文", 2)
    add_para(doc, "基于 Web 的管理系统应当将表现层关注点与业务规则、数据持久化相分离。在分层架构中，客户端负责用户交互，控制层接收请求，业务层执行业务规则，持久化层保存结构化数据。这种分离能够提升系统可维护性，因为某一层的修改通常只会对其他层产生有限影响。")
    add_para(doc, "对于预约类系统而言，一致性比简单插入记录更加重要。应用程序在确认预约之前，必须检查资源可用性、时间冲突、用户权限和支付状态。当多个用户并发操作时，需要通过锁机制或数据库约束保护员工排班、商品库存等共享资源。")
    add_para(doc, "安全性也是现代 Web 应用的基本要求。密码应以加密形式保存，请求应经过身份认证，用户只能访问与自身角色匹配的资源。对于前后端分离系统，基于令牌的认证方式被广泛使用，它能够保持服务端无状态，并简化跨平台访问。")
    page_break(doc)

    add_heading(doc, "附录  程序源代码", 1)
    add_heading(doc, "模块的核心代码", 2)
    add_code(doc, "附录程序1 预约状态流转规则", """
private static final Map<String, Set<String>> FLOW = Map.of(
        "PENDING", Set.of("APPROVED", "REJECTED"),
        "APPROVED", Set.of("WAIT_SERVICE"),
        "WAIT_SERVICE", Set.of("IN_SERVICE"),
        "IN_SERVICE", Set.of("COMPLETED")
);
""")
    add_code(doc, "附录程序2 技师推荐评分", """
long dayLoad = appointmentMapper.countDayLoad(technician.getId(), appointmentTime.toLocalDate());
double rating = technician.getScore() == null ? 5.0D : technician.getScore().doubleValue();
double score = 60 + rating * 8 - dayLoad * 10;
""")
    add_code(doc, "附录程序3 统一响应结构", """
public class ApiResponse<T> {
    private Integer code;
    private String message;
    private T data;
    public static <T> ApiResponse<T> success(T data) {
        return new ApiResponse<>(0, "success", data);
    }
}
""")
    add_code(doc, "附录程序4 支付业务类型", """
paymentOrder.setBusinessType("PRODUCT_ORDER");   // 商品订单支付
paymentOrder.setBusinessType("APPOINTMENT");     // 服务预约支付
paymentOrder.setBusinessType("WALLET_RECHARGE"); // 钱包充值
""")
    page_break(doc)

    add_heading(doc, "致  谢", 1)
    add_para(doc, "本次毕业设计从选题、需求分析、系统设计、编码实现到测试整理，得到了指导教师张腾老师的悉心指导。老师在课题方向、功能范围、论文结构和系统实现细节方面给予了耐心帮助，使我能够逐步完善系统功能并形成较完整的毕业设计说明书，在此表示诚挚感谢。")
    add_para(doc, "同时感谢学院老师在专业课程学习中给予的指导，感谢同学在系统测试、问题讨论和论文修改过程中提供的帮助。通过本次毕业设计，我进一步理解了 Spring Boot、Vue、MySQL、Redis 等技术在实际管理系统中的综合应用，也提升了需求分析、数据库设计、业务建模、前后端联调和文档整理能力。")


def main():
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    set_doc_defaults(doc)
    configure_footer(doc)
    add_cover(doc)
    add_abstract(doc)
    add_toc(doc)
    add_chapter1(doc)
    add_chapter2(doc)
    add_chapter3(doc)
    add_chapter4(doc)
    add_chapter5(doc)
    add_conclusion_refs_translation_appendix(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
